import os
import sys
import json
import ipaddress
import pandas as pd
from openpyxl import load_workbook

SSH_DOMAIN = "lab.local"
FLOW_POLICY_SHEET = "FLOW_POLICY"
DEFAULT_ETHERCHANNEL_PORTS = 1
DEFAULT_ISP_VLAN = 200
# Set per variant below. Explicit Excel value always wins.
DEFAULT_SKIPPED_PORTS = 0
DEFAULT_RSPAN_VLAN = 900
VALID_SPAN_MODES = {"SPAN", "RSPAN", "ERSPAN"}
MONITORING_VRF = "MONITORING"
DEFAULT_ERSPAN_ID = 100
REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN = True
RADIUS_AUTH_PORT = 1812
RADIUS_ACCT_PORT = 1813


def _clean_header(value):
    if pd.isna(value):
        return ""
    return str(value).strip()



def _site_sheet_names(file):
    return [
        name
        for name in load_workbook(file, read_only=True).sheetnames
        if str(name).strip().upper() != FLOW_POLICY_SHEET
    ]


def _extract_table_blocks(df):
    """Return contiguous non-empty blocks as DataFrames with first row as header."""
    blocks = []
    start = None
    non_empty = ~df.isna().all(axis=1)

    for idx, has_data in non_empty.items():
        if has_data and start is None:
            start = idx
        elif not has_data and start is not None:
            raw = df.loc[start:idx - 1].copy()
            raw = raw.dropna(axis=1, how="all")
            if not raw.empty:
                headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
                raw.columns = headers
                table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
                blocks.append((set(headers), table))
            start = None

    if start is not None:
        raw = df.loc[start:].copy().dropna(axis=1, how="all")
        if not raw.empty:
            headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
            raw.columns = headers
            table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
            blocks.append((set(headers), table))

    return blocks


def read_sheet(filename, sheet):
    """
    Reads the switch-related tables by their headers instead of fixed row numbers.
    This makes the sheet resilient to inserted/removed blank rows and extra columns.
    """
    df = pd.read_excel(filename, sheet_name=sheet, header=None)
    blocks = _extract_table_blocks(df)

    md_top = None
    md_switch = None
    swi_data = None
    ip_data = None
    vrf_data = None
    tunnel_data = None

    for headers, table in blocks:
        if "site" in headers and "router-id" in headers:
            md_top = table
        elif "site" in headers and "secret" in headers and "router-id" not in headers:
            md_switch = table
        elif "SW" in headers:
            swi_data = table
        elif {"vrf", "vlan", "mask", "address min"}.issubset(headers):
            ip_data = table
        elif {"vrf", "loopback", "laddr"}.issubset(headers):
            vrf_data = table
        elif {"tunnel id", "vrf"}.issubset(headers):
            tunnel_data = table

    missing = []
    if md_top is None:
        missing.append("top metadata table (site/router-id)")
    if md_switch is None:
        missing.append("switch metadata table (site/secret)")
    if swi_data is None:
        missing.append("switch data table (SW/...)")

    if missing:
        raise ValueError(
            f"Kunne ikke finne forventede tabeller i ark '{sheet}': {', '.join(missing)}"
        )

    return {
        "md": md_switch,
        "swi_data": swi_data,
        "md_top": md_top,
        "ip_data": ip_data,
        "vrf_data": vrf_data,
        "tunnel_data": tunnel_data,
    }


def _value(obj, names, default=None):
    for name in names:
        if name in obj.index:
            value = obj.get(name)
            if not pd.isna(value) and str(value).strip() != "":
                return value
    return default


def _int_value(obj, names, default=0):
    value = _value(obj, names, default)
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return int(default)


def _normalise_compat_choice(value, field, allowed, default):
    if value is None or pd.isna(value) or str(value).strip() == "":
        return default
    text = str(value).strip().upper().replace("-", "_").replace(" ", "_")
    aliases = {
        "OLD": "OLD", "LEGACY": "LEGACY", "GAMMEL": "OLD",
        "NEW": "NEW", "NY": "NEW",
        "AUTHENTICATION": "AUTH", "AUTH": "AUTH",
        "ACCESSSESSION": "ACCESS_SESSION", "ACCESS_SESSION": "ACCESS_SESSION",
        "AUTO": "AUTO", "DOT1Q": "DOT1Q_CMD", "DOT1Q_CMD": "DOT1Q_CMD",
        "WITH_CMD": "DOT1Q_CMD", "NO_CMD": "NO_CMD", "NONE": "NO_CMD",
        "5": "5", "8": "8", "9": "9",
    }
    text = aliases.get(text, text)
    if text not in allowed:
        raise ValueError(
            f"Ugyldig {field}='{value}'. Gyldige verdier: {', '.join(allowed)}"
        )
    return text


def _top_compat(md_top, field, allowed, default):
    if md_top is None or md_top.empty:
        return default
    value = _value(md_top.iloc[0], [field, field.replace("_", " ")], None)
    return _normalise_compat_choice(value, field, allowed, default)


def _switch_compat(row, md_top, field, allowed, default):
    # Per-switch cell wins; blank inherits the site's top metadata default.
    value = _value(row, [field, field.replace("_", " ")], None)
    if value is None:
        return _top_compat(md_top, field, allowed, default)
    return _normalise_compat_choice(value, field, allowed, default)


def _secret_command(prefix, value, secret_type):
    value = "" if pd.isna(value) else str(value).strip()
    if not value:
        raise ValueError(f"Tom secret/passordverdi for '{prefix}'.")
    detected = None
    for stype in ("5", "8", "9"):
        if value.startswith(f"${stype}$"):
            detected = stype
            break
    if secret_type == "AUTO":
        if detected:
            return f"{prefix} secret {detected} {value}"
        return f"{prefix} secret 0 {value}"
    if detected and detected != secret_type:
        raise ValueError(
            f"secret_type={secret_type}, men verdien for '{prefix}' er et type {detected}-hash. "
            "Bytt enten secret_type eller bruk en hash av riktig type."
        )
    if detected is None:
        raise ValueError(
            f"secret_type={secret_type} krever en ferdig type {secret_type}-hash for '{prefix}'. "
            "Bruk AUTO dersom cellen inneholder plaintext."
        )
    return f"{prefix} secret {secret_type} {value}"


def _dot1x_port_control_command(dot1x_syntax):
    return {
        "LEGACY": "dot1x port-control auto",
        "AUTH": "authentication port-control auto",
        "ACCESS_SESSION": "access-session port-control auto",
    }[dot1x_syntax]


def _include_trunk_encapsulation_command(mode):
    # AUTO intentionally preserves the legacy generator behaviour. Platforms
    # with fixed 802.1Q encapsulation should explicitly select NO_CMD.
    return mode in {"AUTO", "DOT1Q_CMD"}


def _site_id(value):
    try:
        f = float(value)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _switch_id(value):
    return _site_id(value)


def _parse_vlan_allocation(value):
    """
    Parse format such as: 10.1-20.5-30.5-40.2
    Returns [(10,1), (20,5), ...]. VLANs with zero access ports are preserved.
    """
    if pd.isna(value) or str(value).strip() == "":
        return []

    vlan_info = []
    for part in str(value).strip().split("-"):
        part = part.strip()
        if not part:
            continue
        if "." not in part:
            raise ValueError(
                f"Ugyldig vlan-antall '{value}'. Forventet f.eks. '10.1-20.5'."
            )
        vlan_s, count_s = part.split(".", 1)
        vlan_info.append((int(float(vlan_s)), int(float(count_s))))
    return vlan_info


def _ordered_unique(values):
    seen = set()
    result = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _interface_key(prefix, ports):
    ports = list(ports)
    if not ports:
        raise ValueError("Tom portliste kan ikke gjøres om til interface-kommando")
    if len(ports) == 1:
        return f"interface {prefix}{ports[0]}"
    return f"interface range {prefix}{ports[0]} - {ports[-1]}"


def _source_interface_ranges(ports, intf_prefix, direction="rx"):
    """Compress physical interface numbers into Cisco monitor-session ranges.

    Examples:
      [1]       -> source interface g0/1 rx
      [1,2,3]   -> source interface g0/1 - 3 rx
      [1,2,5,6] -> source interface g0/1 - 2 rx
                   source interface g0/5 - 6 rx
    """
    ports = sorted(set(int(p) for p in ports))
    if not ports:
        return []

    lines = []
    run_start = ports[0]
    run_end = ports[0]

    def emit(start, end):
        if start == end:
            return f"source interface {intf_prefix}{start} {direction}"
        return f"source interface {intf_prefix}{start} - {end} {direction}"

    for port in ports[1:]:
        if port == run_end + 1:
            run_end = port
            continue

        lines.append(emit(run_start, run_end))
        run_start = run_end = port

    lines.append(emit(run_start, run_end))
    return lines


def _erspan_source_interface_lines(plan, intf_prefix):
    """Build low-duplication ERSPAN sources for one switch.

    - Client traffic is mirrored only when it ENTERS a client access port (rx).
      The same frame is therefore not mirrored again while traversing
      switch-to-switch trunks.
    - Contiguous client access ports are emitted as a Cisco source-interface
      range to keep the generated configuration compact.
    - On SW1 at a spoke site, the router-facing uplink is also mirrored rx.
      This captures traffic entering the LAN from the router/WAN, so Suricata
      still sees both directions of routed client flows.
    - MGMT/server ports, monitoring/sensor ports, downlinks, unused ports and
      downstream-switch uplinks are deliberately excluded.
    """
    lines = []

    # Access ports are allocated contiguously in the current port model, but
    # use the generic range helper so the output also stays correct if gaps are
    # introduced later.
    access_start = plan["first_access_port"]
    access_ports = list(
        range(access_start, access_start + plan["allocated_access_ports"])
    )
    lines.extend(_source_interface_ranges(access_ports, intf_prefix, "rx"))

    # SW1 has exactly one physical router-facing uplink in this architecture.
    # Keep it separate from the access-port range even if interface numbers
    # happen to be adjacent: it has a different monitoring purpose.
    if plan["is_primary_switch"]:
        lines.extend(_source_interface_ranges(plan["uplink_ports"], intf_prefix, "rx"))

    return lines


def _get_isp_vlan(md):
    if not md.empty:
        value = _value(md.iloc[0], ["ISP-VLAN", "isp_vlan", "ISP VLAN"], None)
        if value is not None:
            return int(float(value))
    # Product architecture default; only matters on a switch that actually has VLAN 200.
    return DEFAULT_ISP_VLAN


def _get_etherchannel_ports(md):
    if md.empty:
        return DEFAULT_ETHERCHANNEL_PORTS
    value = _int_value(
        md.iloc[0],
        ["etherchan_num", "etherchannel_num", "etherchannel_ports", "ports_per_channel"],
        DEFAULT_ETHERCHANNEL_PORTS,
    )
    if value < 1:
        raise ValueError("etherchan_num må være minst 1")
    return value


def _is_true(value):
    """Godta True/TRUE/1/yes/ja/x fra sheets."""
    if pd.isna(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    return str(value).strip().lower() in {"true", "1", "yes", "ja", "y", "x"}


def _get_span_mode(md):
    """
    Returner valgt speilingsmodus for sitet.

    Gyldige verdier i Excel:
      - SPAN
      - RSPAN
      - ERSPAN
      - tom / NONE / OFF / INGEN = deaktivert
    """
    if md.empty:
        return None

    value = _value(md.iloc[0], ["span_mode", "span mode", "spanmode"], None)
    if value is None:
        return None

    mode = str(value).strip().upper()
    if mode in {"", "NONE", "OFF", "INGEN", "DISABLED", "FALSE", "0"}:
        return None

    if mode not in VALID_SPAN_MODES:
        raise ValueError(
            f"Ugyldig span_mode '{value}'. Gyldige verdier er "
            f"{', '.join(sorted(VALID_SPAN_MODES))}, eller tom celle for ingen speiling."
        )
    return mode


def _get_rspan_vlan(md):
    if md.empty:
        return DEFAULT_RSPAN_VLAN
    vlan = _int_value(md.iloc[0], ["rspan_vlan", "rspan vlan"], DEFAULT_RSPAN_VLAN)
    if not 1 <= vlan <= 4094 or vlan == 999:
        raise ValueError(f"Ugyldig RSPAN-VLAN {vlan}. Velg VLAN 1-4094, men ikke 999.")
    return vlan


def _get_management_server_ip(md_top, names, label):
    """Read and validate a management-service IP from the top Excel metadata."""
    if md_top is None or md_top.empty:
        raise ValueError(f"{label}-server mangler i Excel-metadata.")

    value = _value(md_top.iloc[0], names, None)
    if value is None:
        raise ValueError(
            f"{label}-server mangler i Excel. Forventet felt: {names[0]}."
        )

    try:
        return str(ipaddress.ip_address(str(value).strip()))
    except ValueError as exc:
        raise ValueError(f"Ugyldig {label}-server-IP i Excel: {value}") from exc


def _get_tacacs_server_ip(md_top):
    return _get_management_server_ip(
        md_top,
        ["tacacs_server_ip", "tacacs server ip", "tacacs_server"],
        "TACACS",
    )


def _get_syslog_server_ip(md_top):
    return _get_management_server_ip(
        md_top,
        ["syslog_server_ip", "rsyslog_server_ip", "syslog server ip", "syslog_server"],
        "Syslog",
    )


def _get_radius_server_ip(md_top, required=False):
    """Read an optional central RADIUS server address from top metadata."""
    if md_top is None or md_top.empty:
        if required:
            raise ValueError("RADIUS-server mangler i Excel-metadata.")
        return None

    value = _value(
        md_top.iloc[0],
        ["radius_server_ip", "radius server ip", "radius_server"],
        None,
    )
    if value is None:
        if required:
            raise ValueError("RADIUS-server mangler i Excel. Forventet felt: radius_server_ip.")
        return None
    try:
        return str(ipaddress.ip_address(str(value).strip()))
    except ValueError as exc:
        raise ValueError(f"Ugyldig RADIUS-server-IP i Excel: {value}") from exc


def _get_radius_key(md_top, required=False):
    if md_top is None or md_top.empty:
        if required:
            raise ValueError("RADIUS-key mangler i Excel-metadata.")
        return None
    value = _value(md_top.iloc[0], ["radius_key", "radius key"], None)
    if value is None or not str(value).strip():
        if required:
            raise ValueError("RADIUS-key mangler i Excel. Forventet felt: radius_key.")
        return None
    return str(value).strip()


def _get_snmpv3_settings(md_top):
    """Return validated SNMPv3 polling settings, or None when disabled."""
    if md_top is None or md_top.empty:
        return None

    row = md_top.iloc[0]
    enabled = _value(
        row,
        ["snmpv3_enabled", "snmpv3 enabled", "snmp_enabled", "snmp enabled"],
        False,
    )
    if not _is_true(enabled):
        return None

    server_ip = _get_management_server_ip(
        md_top,
        ["snmp_server_ip", "snmp server ip", "nms_server_ip", "nms server ip"],
        "SNMP/NMS",
    )
    user = _value(row, ["snmp_user", "snmp user"], None)
    auth_password = _value(row, ["snmp_auth_password", "snmp auth password"], None)
    priv_password = _value(row, ["snmp_priv_password", "snmp priv password"], None)

    missing = []
    if user is None or not str(user).strip():
        missing.append("snmp_user")
    if auth_password is None or not str(auth_password).strip():
        missing.append("snmp_auth_password")
    if priv_password is None or not str(priv_password).strip():
        missing.append("snmp_priv_password")
    if missing:
        raise ValueError(
            "SNMPv3 er aktivert, men følgende Excel-felt mangler: " + ", ".join(missing)
        )

    user = str(user).strip()
    auth_password = str(auth_password).strip()
    priv_password = str(priv_password).strip()
    if any(ch.isspace() for ch in user):
        raise ValueError("snmp_user kan ikke inneholde mellomrom")
    if len(auth_password) < 8:
        raise ValueError("snmp_auth_password må være minst 8 tegn")
    if len(priv_password) < 8:
        raise ValueError("snmp_priv_password må være minst 8 tegn")

    return {
        "server_ip": server_ip,
        "user": user,
        "auth_password": auth_password,
        "priv_password": priv_password,
    }


def create_snmpv3_config(md_top):
    """Generate read-only SNMPv3 authPriv polling config restricted to the NMS."""
    my_data = {"config": {}, "network_info": {}}
    settings = _get_snmpv3_settings(md_top)
    if settings is None:
        return my_data

    server_ip = settings["server_ip"]
    user = settings["user"]
    auth_password = settings["auth_password"]
    priv_password = settings["priv_password"]

    my_data["config"]["ip access-list standard SNMP-NMS-ONLY"] = [
        f"permit host {server_ip}",
        "deny any",
        "exit",
    ]
    my_data["config"]["snmp-server view NMS-READ iso included"] = []
    my_data["config"]["snmp-server group NMS v3 priv read NMS-READ access SNMP-NMS-ONLY"] = []
    my_data["config"][
        f"snmp-server user {user} NMS v3 auth sha {auth_password} priv aes 128 {priv_password}"
    ] = []
    return my_data


def _get_dot1x_vlans(md):
    """Return VLANs where user/client access ports require IEEE 802.1X."""
    if md is None or md.empty:
        return set()
    value = _value(md.iloc[0], ["dot1x_vlans", "dot1x vlans", "8021x_vlans"], None)
    if value is None:
        return set()

    raw = str(value).replace(";", ",").replace(" ", ",")
    result = set()
    for token in raw.split(","):
        token = token.strip()
        if not token:
            continue
        try:
            vlan = int(float(token))
        except ValueError as exc:
            raise ValueError(f"Ugyldig VLAN i dot1x_vlans: {token}") from exc
        if not 1 <= vlan <= 4094 or vlan == 999:
            raise ValueError(f"Ugyldig VLAN i dot1x_vlans: {vlan}")
        result.add(vlan)
    return result


def _find_vrf_rows(df, vrf_name):
    if df is None or df.empty or "vrf" not in df.columns:
        return df.iloc[0:0] if df is not None else pd.DataFrame()
    names = df["vrf"].astype(str).str.strip().str.upper()
    return df[names == vrf_name.upper()]


def _get_monitoring_service(ip_data, vrf_data, site):
    """Return the MONITORING service definition used by ERSPAN.

    NO-MPLS requires an explicit MONITORING VRF/IP row because it also needs
    Tunnel50/DMVPN transport.  MPLS can derive the fixed reference-architecture
    MONITORING network automatically when those rows are absent:
      Site N -> VLAN 50, 10.(50+N).0.0/24, gateway .1.
    """
    vrf_rows = _find_vrf_rows(vrf_data, MONITORING_VRF)
    ip_rows = _find_vrf_rows(ip_data, MONITORING_VRF)

    if vrf_rows.empty or ip_rows.empty:
        if not REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN:
            try:
                site_no = int(float(site))
                second_octet = 50 + site_no
                if not 0 <= second_octet <= 255:
                    raise ValueError
                network = ipaddress.ip_network(f"10.{second_octet}.0.0/24")
                return {
                    "vrf": MONITORING_VRF,
                    "vlan": 50,
                    "mask": "255.255.255.0",
                    "gateway": str(network.network_address + 1),
                    "network": network,
                }
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Site {site}: kunne ikke utlede MONITORING-nett automatisk for MPLS."
                ) from exc

        if vrf_rows.empty:
            raise ValueError(
                f"Site {site}: span_mode=ERSPAN krever VRF '{MONITORING_VRF}' "
                "i Excel sin VRF-tabell for NO-MPLS router-transport."
            )
        raise ValueError(
            f"Site {site}: span_mode=ERSPAN krever en egen VLAN/subnett-rad med "
            f"vrf={MONITORING_VRF} i IP-tabellen."
        )

    if len(vrf_rows) != 1 or len(ip_rows) != 1:
        raise ValueError(
            f"Site {site}: ERSPAN forventer nøyaktig en {MONITORING_VRF}-rad i "
            "VRF-tabellen og en i IP/VLAN-tabellen."
        )

    row = ip_rows.iloc[0]
    try:
        vlan = int(float(row["vlan"]))
        mask = str(row["mask"]).strip()
        gateway = str(row.get("gateway", row.get("address min", ""))).strip()
        network_value = row.get("nett id", gateway)
        network = ipaddress.ip_network(f"{network_value}/{mask}", strict=False)
        gateway_ip = ipaddress.ip_address(gateway)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Site {site}: ugyldig MONITORING VLAN/subnett/gateway i Excel.") from exc

    if not 1 <= vlan <= 4094 or vlan == 999:
        raise ValueError(f"Site {site}: ugyldig MONITORING VLAN {vlan}.")
    if gateway_ip not in network:
        raise ValueError(
            f"Site {site}: MONITORING gateway {gateway} ligger ikke i {network}."
        )

    return {
        "vrf": MONITORING_VRF,
        "vlan": vlan,
        "mask": mask,
        "gateway": gateway,
        "network": network,
    }


def _get_switch_monitoring_ip(row, monitoring, site, sw_id):
    value = _value(
        row,
        ["MONITORING ip", "monitoring ip", "MONITORING IP", "monitoring_ip"],
        None,
    )
    if value is None:
        raise ValueError(
            f"Site {site} SW{sw_id}: span_mode=ERSPAN krever en egen 'MONITORING ip' "
            "i switch-tabellen. MGMT ip kan ikke brukes som ERSPAN origin."
        )

    try:
        address = ipaddress.ip_address(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"Site {site} SW{sw_id}: ugyldig MONITORING ip: {value}") from exc

    network = monitoring["network"]
    if address not in network or address in {network.network_address, network.broadcast_address}:
        raise ValueError(
            f"Site {site} SW{sw_id}: MONITORING ip {address} er ikke en gyldig host i {network}."
        )
    if str(address) == monitoring["gateway"]:
        raise ValueError(
            f"Site {site} SW{sw_id}: MONITORING ip kan ikke være gateway {monitoring['gateway']}."
        )
    return str(address)


def _validate_erspan_architecture_for_workbook(file, sheets):
    """Validate ERSPAN and derive the destination from HUB SW1 MONITORING IP.

    The user does not enter an ERSPAN destination manually.  In ERSPAN mode,
    HUB SW1 terminates the GRE/ERSPAN stream and forwards the decapsulated
    mirrored frames out its dedicated physical monitor destination port.
    """
    records = []
    hubs = []
    active_erspan_sites = set()

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        md_top = sheet_data["md_top"]
        site = _site_id(md.iloc[0]["site"])
        ip_data = sheet_data.get("ip_data")
        vrf_data = sheet_data.get("vrf_data")
        tunnel_data = sheet_data.get("tunnel_data")
        mode = _get_span_mode(md)
        is_hub = _is_hub_site(md, md_top, site)

        record = {
            "sheet": sheet,
            "site": site,
            "data": sheet_data,
            "mode": mode,
            "is_hub": is_hub,
        }
        records.append(record)
        if is_hub:
            hubs.append(record)
        if mode == "ERSPAN":
            active_erspan_sites.add(site)

    if not active_erspan_sites:
        return {
            "active_sites": set(),
            "hub_site": None,
            "destination_ip": None,
            "erspan_id": DEFAULT_ERSPAN_ID,
        }

    if len(hubs) != 1:
        raise ValueError(
            f"ERSPAN krever nøyaktig ett HUB-site. Fant {len(hubs)} HUB-sites."
        )

    hub = hubs[0]
    if hub["mode"] != "ERSPAN":
        raise ValueError(
            f"ERSPAN er aktivert på site(s) {sorted(active_erspan_sites)}, men HUB Site "
            f"{hub['site']} har span_mode={hub['mode'] or 'OFF'}. HUB må også stå i ERSPAN."
        )

    hub_data = hub["data"]
    hub_monitoring = _get_monitoring_service(
        hub_data.get("ip_data"), hub_data.get("vrf_data"), hub["site"]
    )
    hub_sw1_rows = hub_data["swi_data"][
        hub_data["swi_data"]["SW"].apply(_switch_id) == "1"
    ]
    if len(hub_sw1_rows) != 1:
        raise ValueError(
            f"HUB Site {hub['site']}: ERSPAN krever nøyaktig en SW1 i switch-tabellen."
        )
    destination = _get_switch_monitoring_ip(
        hub_sw1_rows.iloc[0], hub_monitoring, hub["site"], "1"
    )

    # Destination must be the HUB SW1 MONITORING address and never a MGMT address.
    mgmt_rows = _find_vrf_rows(hub_data.get("ip_data"), "MGMT")
    if not mgmt_rows.empty:
        r = mgmt_rows.iloc[0]
        mgmt_net = ipaddress.ip_network(
            f"{r.get('nett id', r.get('address min'))}/{r['mask']}", strict=False
        )
        if ipaddress.ip_address(destination) in mgmt_net:
            raise ValueError(
                f"HUB SW1 ERSPAN-destinasjon {destination} ligger i MGMT-nettet {mgmt_net}."
            )

    # Validate every active site's MONITORING transport and unique switch origin IPs.
    globally_seen_monitoring_ips = set()
    for record in records:
        if record["mode"] != "ERSPAN":
            continue

        site = record["site"]
        sheet_data = record["data"]
        monitoring = _get_monitoring_service(
            sheet_data.get("ip_data"), sheet_data.get("vrf_data"), site
        )

        for _, sw_row in sheet_data["swi_data"].iterrows():
            sw_id = _switch_id(sw_row["SW"])
            mon_ip = _get_switch_monitoring_ip(sw_row, monitoring, site, sw_id)
            if mon_ip in globally_seen_monitoring_ips:
                raise ValueError(
                    f"MONITORING ip {mon_ip} er brukt av flere switcher i arbeidsboken."
                )
            globally_seen_monitoring_ips.add(mon_ip)

        if REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN:
            tun_rows = _find_vrf_rows(sheet_data.get("tunnel_data"), MONITORING_VRF)
            if tun_rows.empty:
                raise ValueError(
                    f"Site {site}: NO-MPLS + ERSPAN krever en egen DMVPN/EIGRP-tunnel "
                    f"for VRF {MONITORING_VRF}. Legg til en MONITORING-rad i tunnel-tabellen."
                )

    return {
        "active_sites": active_erspan_sites,
        "hub_site": hub["site"],
        "destination_ip": destination,
        "erspan_id": DEFAULT_ERSPAN_ID,
    }


def _validate_management_servers_for_workbook(file, sheets):
    """Validate centralized TACACS/Syslog/RADIUS and switch MGMT identities."""
    values = []
    hub_record = None
    switch_mgmt_ips = {}

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        md_top = sheet_data["md_top"]
        site = _site_id(md.iloc[0]["site"])
        tacacs = _get_tacacs_server_ip(md_top)
        syslog = _get_syslog_server_ip(md_top)
        radius = _get_radius_server_ip(md_top, required=False)
        radius_key = _get_radius_key(md_top, required=False)
        dot1x_vlans = _get_dot1x_vlans(md)

        if (radius is None) != (radius_key is None):
            raise ValueError(
                f"Site {site}: radius_server_ip og radius_key må enten begge være satt eller begge være tomme."
            )
        if dot1x_vlans and radius is None:
            raise ValueError(
                f"Site {site}: dot1x_vlans er satt, men radius_server_ip/radius_key mangler."
            )

        values.append((site, tacacs, syslog, radius, radius_key))
        if _is_hub_site(md, md_top, site):
            if hub_record is not None:
                raise ValueError("Flere HUB-sites funnet ved validering av management-servere.")
            hub_record = (site, sheet_data, tacacs, syslog, radius, radius_key)

        # Validate every switch MGMT IP against this site's MGMT subnet. These
        # addresses are also the NAS/client addresses written to FreeRADIUS.
        mgmt_rows = _find_vrf_rows(sheet_data.get("ip_data"), "MGMT")
        if mgmt_rows.empty:
            raise ValueError(f"Site {site}: finner ikke MGMT-subnett for switch-validering.")
        r = mgmt_rows.iloc[0]
        mgmt_net = ipaddress.ip_network(
            f"{r.get('nett id', r.get('address min'))}/{r['mask']}", strict=False
        )
        for _, sw_row in sheet_data["swi_data"].iterrows():
            sw_id = _switch_id(sw_row["SW"])
            raw_ip = _value(sw_row, ["MGMT ip", "mgmt ip", "mgmt_ip"], None)
            if raw_ip is None:
                raise ValueError(f"Site {site} SW{sw_id}: MGMT ip mangler.")
            try:
                addr = ipaddress.ip_address(str(raw_ip).strip())
            except ValueError as exc:
                raise ValueError(f"Site {site} SW{sw_id}: ugyldig MGMT ip {raw_ip}.") from exc
            if addr not in mgmt_net or addr in {mgmt_net.network_address, mgmt_net.broadcast_address}:
                raise ValueError(
                    f"Site {site} SW{sw_id}: MGMT ip {addr} må være gyldig host i {mgmt_net}."
                )
            if str(addr) in switch_mgmt_ips:
                raise ValueError(
                    f"MGMT ip {addr} brukes både av {switch_mgmt_ips[str(addr)]} og Site {site} SW{sw_id}."
                )
            switch_mgmt_ips[str(addr)] = f"Site {site} SW{sw_id}"

    if hub_record is None:
        raise ValueError("Fant ikke HUB-site ved validering av management-servere.")

    tacacs_values = {x[1] for x in values}
    syslog_values = {x[2] for x in values}
    radius_values = {x[3] for x in values}
    radius_keys = {x[4] for x in values}
    if len(tacacs_values) != 1 or len(syslog_values) != 1 or len(radius_values) != 1 or len(radius_keys) != 1:
        detail = ", ".join(
            f"Site {site}: TACACS={tacacs}, Syslog={syslog}, RADIUS={radius}"
            for site, tacacs, syslog, radius, _radius_key in values
        )
        raise ValueError(
            "TACACS/Syslog/RADIUS-innstillingene må være konsistente mellom site-arkene. " + detail
        )

    site, sheet_data, tacacs, syslog, radius, radius_key = hub_record
    mgmt_rows = _find_vrf_rows(sheet_data.get("ip_data"), "MGMT")
    if mgmt_rows.empty:
        raise ValueError(f"HUB Site {site}: finner ikke MGMT-subnett for servervalidering.")
    r = mgmt_rows.iloc[0]
    mgmt_net = ipaddress.ip_network(
        f"{r.get('nett id', r.get('address min'))}/{r['mask']}", strict=False
    )
    server_pairs = [("TACACS", tacacs), ("Syslog", syslog)]
    if radius is not None:
        server_pairs.append(("RADIUS", radius))
    for label, server in server_pairs:
        addr = ipaddress.ip_address(server)
        if addr not in mgmt_net or addr in {mgmt_net.network_address, mgmt_net.broadcast_address}:
            raise ValueError(
                f"{label}-server {server} må ligge som gyldig host i HUB MGMT-nettet {mgmt_net}."
            )

    return {
        "tacacs": tacacs,
        "syslog": syslog,
        "radius": radius,
        "radius_key": radius_key,
        "switch_mgmt_ips": switch_mgmt_ips,
        "same_server": tacacs == syslog == radius if radius is not None else tacacs == syslog,
    }


def _validate_span_modes_for_workbook(file, sheets):
    """
    Alle sites som har speiling aktivert må bruke samme modus.
    Tom span_mode betyr at speiling er deaktivert på det sitet og regnes
    derfor ikke som en konflikt.
    """
    modes_by_site = {}
    enabled_modes = set()

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        site = _site_id(md.iloc[0]["site"])
        mode = _get_span_mode(md)
        modes_by_site[site] = mode
        if mode is not None:
            enabled_modes.add(mode)

    if len(enabled_modes) > 1:
        details = ", ".join(
            f"Site {site}={mode or 'OFF'}"
            for site, mode in sorted(modes_by_site.items())
        )
        raise ValueError(
            "Konflikt i span_mode mellom sites. Aktive sites må bruke samme "
            f"speilingsmodus. Fant: {details}"
        )

    return modes_by_site


def _is_hub_site(md, md_top, site):
    """
    Determine whether this sheet represents the hub site.

    Optional Excel overrides supported:
      - is_hub / hub = TRUE/FALSE
      - role / site_role = HUB
      - hub_site = site number

    Backward-compatible product default: site 1 is the hub.
    """
    site = _site_id(site)

    for table in (md, md_top):
        if table is None or table.empty:
            continue
        row = table.iloc[0]

        explicit = _value(row, ["is_hub", "hub", "HUB"], None)
        if explicit is not None:
            return _is_true(explicit)

        role = _value(row, ["role", "site_role", "site role"], None)
        if role is not None:
            return str(role).strip().lower() == "hub"

        hub_site = _value(row, ["hub_site", "hub site"], None)
        if hub_site is not None:
            return _site_id(hub_site) == site

    return site == "1"


def _effective_vlan_info(row, md, is_hub):
    """
    VLAN 200 is the local ISP transit VLAN in this architecture.
    It is only valid on SW1 at the hub and must never be extended to
    downstream switches or spoke sites.
    """
    vlan_info = _parse_vlan_allocation(row.get("vlan-antall", ""))
    isp_vlan = _get_isp_vlan(md)
    sw_id = _switch_id(row.get("SW"))

    if not (is_hub and sw_id == "1"):
        vlan_info = [(vlan, count) for vlan, count in vlan_info if vlan != isp_vlan]

    return vlan_info


def _site_service_vlans(swi_data, md, is_hub):
    """Return the union of service VLANs used anywhere in this site.

    vlan-antall controls LOCAL access-port allocation only.  A VLAN that is
    present on SW3 must still exist and be allowed across the transit trunks
    on SW2/SW1 so it can reach the site router.

    VLAN 200 is special: _effective_vlan_info() only keeps it on HUB SW1.
    The trunk generator additionally removes it from switch-to-switch trunks.
    """
    vlans = []
    for _, site_row in swi_data.iterrows():
        for vlan, _count in _effective_vlan_info(site_row, md, is_hub):
            vlans.append(vlan)
    return _ordered_unique(vlans)


def _build_port_plan(row, md, md_top, swi_data, site, is_hub=False):
    """Build deterministic physical port allocation for a switch.

    Every switch keeps one dedicated MGMT access port as before.  On HUB SW1
    that port is the TACACS/RADIUS/Syslog server port when the services share an IP.
    If the service IPs differ, HUB SW1 reserves one additional MGMT port for
    Syslog/Security Onion management.  SPAN/RSPAN reserve a local destination
    port on SW1; ERSPAN reserves a mirror destination port only on HUB SW1.
    """
    has_new_num_ports = (
        "num_ports_tot" in row.index
        and not pd.isna(row.get("num_ports_tot"))
        and str(row.get("num_ports_tot")).strip() != ""
    )

    num_ports = _int_value(row, ["num_ports_tot", "num_ports"], 0)
    skipped_ports = _int_value(
        row,
        ["skiped_ports", "skipped_ports"],
        DEFAULT_SKIPPED_PORTS,
    )
    num_downlinks = _int_value(row, ["num_downlink", "num_downlinks"], 0)
    etherchan_ports = _get_etherchannel_ports(md)
    sw_id = _switch_id(row.get("SW"))
    is_primary_switch = sw_id == "1"

    if num_ports <= 0:
        raise ValueError(f"SW{sw_id}: num_ports må være > 0")
    if skipped_ports < 0:
        raise ValueError(f"SW{sw_id}: skiped_ports kan ikke være negativ")
    if num_downlinks < 0:
        raise ValueError(f"SW{sw_id}: num_downlink kan ikke være negativ")
    if has_new_num_ports and skipped_ports >= num_ports:
        raise ValueError(
            f"SW{sw_id}: skiped_ports ({skipped_ports}) må være mindre enn "
            f"num_ports_tot ({num_ports})"
        )

    first_usable = skipped_ports
    if has_new_num_ports:
        last_port = num_ports - 1
        usable_count = num_ports - skipped_ports
        skipped_physical_ports = list(range(0, skipped_ports))
    else:
        last_port = skipped_ports + num_ports - 1
        usable_count = num_ports
        skipped_physical_ports = []

    mgmt_port = first_usable
    tacacs_server = _get_tacacs_server_ip(md_top)
    syslog_server = _get_syslog_server_ip(md_top)
    separate_server_ports = (
        is_hub and is_primary_switch and tacacs_server != syslog_server
    )
    syslog_port = mgmt_port + 1 if separate_server_ports else mgmt_port

    span_mode = _get_span_mode(md)
    if span_mode == "RSPAN" and len(swi_data) < 2:
        raise ValueError(
            f"Site {site}: RSPAN krever minst en downstream-switch. "
            "Sitet har bare SW1; bruk SPAN i stedet."
        )

    local_sensor_port_required = (
        is_primary_switch
        and (
            span_mode in {"SPAN", "RSPAN"}
            or (span_mode == "ERSPAN" and is_hub)
        )
    )

    server_port_count = 1 + (1 if separate_server_ports else 0)
    sensor_port = (
        first_usable + server_port_count if local_sensor_port_required else None
    )
    dedicated_ports = server_port_count + (1 if local_sensor_port_required else 0)

    uplink_member_count = 1 if is_primary_switch else etherchan_ports
    downlink_member_count = num_downlinks * etherchan_ports
    trunk_member_count = uplink_member_count + downlink_member_count
    available_access_slots = usable_count - dedicated_ports - trunk_member_count

    if available_access_slots < 0:
        uplink_desc = "1 router-uplink" if is_primary_switch else f"{etherchan_ports}-ports uplink"
        raise ValueError(
            f"SW{sw_id}: ikke nok porter. {usable_count} brukbare porter, men "
            f"dedikerte MGMT/server/sensor-porter + {uplink_desc} + downlinks krever "
            f"{dedicated_ports + trunk_member_count} porter."
        )

    first_access_port = first_usable + dedicated_ports
    last_access_port = first_access_port + available_access_slots - 1

    uplink_start = last_port - uplink_member_count + 1
    uplink_ports = list(range(uplink_start, last_port + 1))

    downlink_groups = []
    first_downlink_top = uplink_start - 1
    for idx in range(num_downlinks):
        group_end = first_downlink_top - idx * etherchan_ports
        group_start = group_end - etherchan_ports + 1
        downlink_groups.append(list(range(group_start, group_end + 1)))

    vlan_info = _effective_vlan_info(row, md, is_hub)
    allocated_access_ports = sum(max(0, count) for _, count in vlan_info)
    if allocated_access_ports > available_access_slots:
        raise ValueError(
            f"SW{sw_id}: vlan-antall bruker {allocated_access_ports} access-porter, "
            f"men bare {available_access_slots} er tilgjengelige etter "
            f"dedikerte porter/uplink/downlinks."
        )

    expected_free = _value(row, ["num_port_ledig"], None)
    if expected_free is not None:
        try:
            expected_free = int(float(expected_free))
            if expected_free != available_access_slots:
                print(
                    f"ADVARSEL SW{sw_id}: num_port_ledig i Excel er {expected_free}, "
                    f"men generatoren beregner {available_access_slots}."
                )
        except (TypeError, ValueError):
            pass

    return {
        "num_ports": num_ports,
        "skipped_ports": skipped_ports,
        "skipped_physical_ports": skipped_physical_ports,
        "mgmt_port": mgmt_port,
        "tacacs_port": mgmt_port,
        "syslog_port": syslog_port,
        "separate_server_ports": separate_server_ports,
        "span_port": sensor_port,
        "sensor_port": sensor_port,
        "span_mode": span_mode,
        "first_access_port": first_access_port,
        "last_access_port": last_access_port,
        "available_access_slots": available_access_slots,
        "allocated_access_ports": allocated_access_ports,
        "etherchan_ports": etherchan_ports,
        "is_primary_switch": is_primary_switch,
        "uplink_is_etherchannel": (not is_primary_switch and etherchan_ports > 1),
        "downlink_is_etherchannel": etherchan_ports > 1 and num_downlinks > 0,
        "num_downlinks": num_downlinks,
        "uplink_ports": uplink_ports,
        "downlink_groups": downlink_groups,
        "vlan_info": vlan_info,
    }


def create_tacacs_config(md_top, switch_row=None):
    my_data = {"config": {}, "network_info": {}}
    tacacs_server = _get_tacacs_server_ip(md_top)
    tacacs_key = md_top.iloc[0].get("tacacs_key", "")

    if pd.isna(tacacs_key) or not str(tacacs_key).strip():
        raise ValueError("TACACS-key mangler i Excel")
    tacacs_key = str(tacacs_key).strip()
    syntax = (
        _switch_compat(switch_row, md_top, "tacacs_syntax", ("OLD", "NEW"), "NEW")
        if switch_row is not None
        else _top_compat(md_top, "tacacs_syntax", ("OLD", "NEW"), "NEW")
    )

    my_data["config"]["aaa new-model"] = []
    if syntax == "NEW":
        my_data["config"]["tacacs server TACACS-SERVER"] = [
            f"address ipv4 {tacacs_server}",
            f"key {tacacs_key}",
            "exit",
        ]
        group_server = "server name TACACS-SERVER"
    else:
        my_data["config"][f"tacacs-server host {tacacs_server} key {tacacs_key}"] = []
        group_server = f"server {tacacs_server}"

    my_data["config"]["aaa group server tacacs+ TACACS-GROUP"] = [
        group_server,
        "ip tacacs source-interface Vlan10",
        "exit",
    ]
    my_data["config"]["aaa authentication login default group TACACS-GROUP local"] = []
    my_data["config"]["aaa authorization exec default group TACACS-GROUP local"] = []
    my_data["config"]["aaa accounting exec default start-stop group TACACS-GROUP"] = []
    my_data["config"]["aaa accounting commands 15 default start-stop group TACACS-GROUP"] = []
    return my_data


def create_radius_dot1x_config(md_top, mgmt_vlan, switch_row=None):
    """Generate switch-side RADIUS + IEEE 802.1X global configuration."""
    my_data = {"config": {}, "network_info": {}}
    radius_server = _get_radius_server_ip(md_top, required=False)
    radius_key = _get_radius_key(md_top, required=False)
    if radius_server is None and radius_key is None:
        return my_data
    if radius_server is None or radius_key is None:
        raise ValueError("radius_server_ip og radius_key må begge være satt for 802.1X.")

    syntax = (
        _switch_compat(switch_row, md_top, "radius_syntax", ("OLD", "NEW"), "NEW")
        if switch_row is not None
        else _top_compat(md_top, "radius_syntax", ("OLD", "NEW"), "NEW")
    )

    if syntax == "NEW":
        my_data["config"]["radius server CLIENT-RADIUS"] = [
            f"address ipv4 {radius_server} auth-port {RADIUS_AUTH_PORT} acct-port {RADIUS_ACCT_PORT}",
            f"key {radius_key}",
            "exit",
        ]
        my_data["config"]["aaa group server radius DOT1X-RADIUS"] = [
            "server name CLIENT-RADIUS",
            "exit",
        ]
        radius_group = "DOT1X-RADIUS"
    else:
        my_data["config"][
            f"radius-server host {radius_server} auth-port {RADIUS_AUTH_PORT} acct-port {RADIUS_ACCT_PORT} key {radius_key}"
        ] = []
        radius_group = "radius"

    my_data["config"][f"ip radius source-interface Vlan{mgmt_vlan}"] = []
    my_data["config"][f"aaa authentication dot1x default group {radius_group}"] = []
    my_data["config"][f"aaa authorization network default group {radius_group}"] = []
    my_data["config"][f"aaa accounting dot1x default start-stop group {radius_group}"] = []
    my_data["config"]["dot1x system-auth-control"] = []
    return my_data


def create_rsyslog_config(md_top):
    my_data = {"config": {}, "network_info": {}}
    rsyslog_server = _get_syslog_server_ip(md_top)

    my_data["config"]["service timestamps log datetime msec show-timezone"] = []
    my_data["config"][f"logging host {rsyslog_server} transport udp port 514"] = []
    my_data["config"]["logging trap informational"] = []
    my_data["config"]["logging buffered 16384 informational"] = []
    my_data["config"]["logging source-interface Vlan10"] = []
    return my_data


def enable_ssh(md, domain=SSH_DOMAIN, mgmt_network=None, mgmt_wildcard=None, secret_type="AUTO"):
    my_data = {"config": {}, "network_info": {}}
    if md.empty:
        return my_data

    row = md.iloc[0]
    username = row.get("brukernavn", "")
    password = row.get("passord", "")
    vty_lines = row.get("vty_lines", "0-4")

    if pd.isna(username) or pd.isna(password):
        return my_data

    username = str(username).strip()
    password = str(password).strip()
    if not username or not password:
        return my_data

    my_data["config"][f"ip domain name {domain}"] = []
    my_data["config"][_secret_command(f"username {username} privilege 15", password, secret_type)] = []
    my_data["config"]["crypto key generate rsa general-keys modulus 2048"] = []
    my_data["config"]["ip ssh version 2"] = []
    my_data["config"]["ip scp server enable"] = []

    vty_cfg = ["login authentication default", "exec-timeout 10 0", "transport input ssh"]
    if mgmt_network and mgmt_wildcard:
        my_data["config"]["ip access-list standard SSH-MGMT-ONLY"] = [
            f"permit {mgmt_network} {mgmt_wildcard}",
            "exit",
        ]
        vty_cfg.insert(0, "access-class SSH-MGMT-ONLY in")
    vty_cfg.append("exit")

    my_data["config"][f"line vty {' '.join(x.strip() for x in str(vty_lines).split('-'))}"] = vty_cfg

    return my_data


def _mgmt_access_port_config(description, mgmt_vlan):
    return [
        f"description {description}",
        "ip arp inspection trust",
        "switchport mode access",
        f"switchport access vlan {mgmt_vlan}",
        "spanning-tree portfast",
        "no shutdown",
        "exit",
    ]


def global_config(md, md_top, swi_data, ip_data, vrf_data, is_hub, erspan_context=None):
    info = {"config": {}, "network_info": {}}
    site = _site_id(md.iloc[0]["site"])
    secret = md.iloc[0].get("secret", "")
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )
    erspan_destination = (
        erspan_context.get("destination_ip")
        if span_mode_site == "ERSPAN" and erspan_context
        else None
    )
    erspan_id = (
        int(erspan_context.get("erspan_id", DEFAULT_ERSPAN_ID))
        if erspan_context
        else DEFAULT_ERSPAN_ID
    )

    if span_mode_site == "ERSPAN" and not erspan_destination:
        raise ValueError("ERSPAN-destinasjon kunne ikke utledes fra HUB SW1 MONITORING ip.")

    tacacs_server = _get_tacacs_server_ip(md_top)
    syslog_server = _get_syslog_server_ip(md_top)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        mgmt_ip = row["MGMT ip"]
        gateway = row["gateway"]
        mask = row["mask"]
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)
        secret_type = _switch_compat(
            row, md_top, "secret_type", ("AUTO", "5", "8", "9"), "AUTO"
        )

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg[f"hostname {sw_name}"] = []
        sw_cfg[_secret_command("enable", secret, secret_type)] = []
        sw_cfg["service tcp-keepalives-in"] = []
        sw_cfg["service tcp-keepalives-out"] = []
        sw_cfg["vtp mode transparent"] = []
        sw_cfg["banner motd ^CKun autorisert tilgang er tillatt. Aktivitet kan bli logget.^C"] = []
        sw_cfg.update(create_tacacs_config(md_top, row)["config"])
        sw_cfg.update(create_radius_dot1x_config(md_top, mgmt_vlan, row)["config"])
        sw_cfg["line console 0"] = [
            "login authentication default",
            "exec-timeout 10 0",
            "logging synchronous",
            "exit",
        ]
        mgmt_network = mgmt_wildcard = None
        try:
            mgmt_net = ipaddress.ip_network(f"{mgmt_ip}/{mask}", strict=False)
            mgmt_network, mgmt_wildcard = str(mgmt_net.network_address), str(mgmt_net.hostmask)
        except (ValueError, TypeError):
            pass
        sw_cfg.update(
            enable_ssh(
                md, mgmt_network=mgmt_network, mgmt_wildcard=mgmt_wildcard,
                secret_type=secret_type
            )["config"]
        )
        sw_cfg.update(create_rsyslog_config(md_top)["config"])
        sw_cfg.update(create_snmpv3_config(md_top)["config"])

        sw_cfg[f"vlan {mgmt_vlan}"] = [f"name MGMT_VLAN_{mgmt_vlan}", "exit"]
        sw_cfg[f"interface vlan {mgmt_vlan}"] = [
            f"ip address {mgmt_ip} {mask}",
            "no shutdown",
            "exit",
        ]

        monitoring_ip = None
        if span_mode_site == "ERSPAN":
            monitoring_ip = _get_switch_monitoring_ip(row, monitoring, site, sw_id)
            monitoring_vlan = monitoring["vlan"]
            monitoring_mask = monitoring["mask"]
            monitoring_gateway = monitoring["gateway"]

            sw_cfg["ip routing"] = []
            sw_cfg[f"vlan {monitoring_vlan}"] = ["name MONITORING", "exit"]
            sw_cfg[f"interface vlan {monitoring_vlan}"] = [
                f"ip address {monitoring_ip} {monitoring_mask}",
                "no shutdown",
                "exit",
            ]
            destination_ip = ipaddress.ip_address(erspan_destination)
            if destination_ip not in monitoring["network"]:
                sw_cfg[
                    f"ip route {erspan_destination} 255.255.255.255 {monitoring_gateway}"
                ] = []

        # Preserve the existing dedicated MGMT port on every switch.  On HUB SW1
        # it becomes the physical server handoff for TACACS/RADIUS/Syslog.
        if is_hub and plan["is_primary_switch"]:
            if tacacs_server == syslog_server:
                mgmt_desc = "Dedicated TACACS/RADIUS/Syslog management server port"
            else:
                mgmt_desc = "Dedicated TACACS management server port"
        else:
            mgmt_desc = f"Dedicated management access port for VLAN {mgmt_vlan}"

        sw_cfg[f"interface {intf_prefix}{plan['tacacs_port']}"] = _mgmt_access_port_config(
            mgmt_desc, mgmt_vlan
        )

        if is_hub and plan["is_primary_switch"] and plan["separate_server_ports"]:
            sw_cfg[f"interface {intf_prefix}{plan['syslog_port']}"] = _mgmt_access_port_config(
                "Dedicated Syslog/Security Onion management server port", mgmt_vlan
            )

        # ERSPAN destination port is a true mirror destination.  It is not an
        # access port and carries no IP/VLAN configuration toward the sensor NIC.
        if (
            span_mode_site == "ERSPAN"
            and is_hub
            and plan["is_primary_switch"]
            and plan["sensor_port"] is not None
        ):
            sw_cfg[f"interface {intf_prefix}{plan['sensor_port']}"] = [
                "description Dedicated Security Onion ERSPAN mirror destination port",
                "no shutdown",
                "exit",
            ]

        span_mode = plan["span_mode"]
        span_vlans = _ordered_unique([mgmt_vlan, *[v for v, _ in plan["vlan_info"]]])
        if span_mode == "ERSPAN" and monitoring is not None:
            span_vlans = [v for v in span_vlans if v != monitoring["vlan"]]

        if span_mode == "SPAN" and plan["span_port"] is not None:
            sw_cfg[f"interface {intf_prefix}{plan['span_port']}"] = [
                "description Dedicated local SPAN destination port for IDS/IPS",
                "no shutdown",
                "exit",
            ]
            sw_cfg[f"monitor session 1 source vlan {' , '.join(map(str, span_vlans))} both"] = []
            sw_cfg[f"monitor session 1 destination interface {intf_prefix}{plan['span_port']}"] = []

        elif span_mode == "RSPAN":
            rspan_vlan = _get_rspan_vlan(md)
            if plan["is_primary_switch"]:
                if plan["span_port"] is None:
                    raise ValueError(
                        f"SW{sw_id}: RSPAN krever en lokal sensorport på primærswitchen."
                    )
                sw_cfg[f"interface {intf_prefix}{plan['span_port']}"] = [
                    "description Dedicated RSPAN destination port for IDS/IPS",
                    "no shutdown",
                    "exit",
                ]
                sw_cfg[f"monitor session 1 source remote vlan {rspan_vlan}"] = []
                sw_cfg[f"monitor session 1 destination interface {intf_prefix}{plan['span_port']}"] = []
            else:
                sw_cfg[f"monitor session 1 source vlan {' , '.join(map(str, span_vlans))} both"] = []
                sw_cfg[f"monitor session 1 destination remote vlan {rspan_vlan}"] = []

        elif span_mode == "ERSPAN":
            if is_hub and plan["is_primary_switch"]:
                # One common ERSPAN-ID lets a single destination session receive
                # mirrored traffic from all remote ERSPAN source switches.
                sw_cfg["monitor session 1 type erspan-destination"] = [
                    "description ERSPAN-TO-SECURITY-ONION",
                    f"destination interface {intf_prefix}{plan['sensor_port']}",
                    "source",
                    f"erspan-id {erspan_id}",
                    f"ip address {erspan_destination}",
                    "no shutdown",
                    "exit",
                    "exit",
                ]
                print(
                    f"ADVARSEL Site {site} SW1: ERSPAN destination-porten kan bare tilhøre en "
                    "monitor-session. Lokal-only trafikk på HUB SW1 speiles derfor ikke i ERSPAN-modus. "
                    "Remote sites speiler klient-ingress og router-uplink-ingress for å redusere duplikater."
                )
            else:
                # Avoid duplicate ERSPAN copies across a multi-switch site:
                # - mirror ingress on real client access ports
                # - on SW1, also mirror ingress from the site router/WAN
                # - never mirror switch-to-switch trunks
                source_lines = _erspan_source_interface_lines(plan, intf_prefix)

                if source_lines:
                    sw_cfg["monitor session 1 type erspan-source"] = [
                        f"description ERSPAN-SITE-{site}-SW{sw_id}",
                        *source_lines,
                        "destination",
                        f"ip address {erspan_destination}",
                        f"erspan-id {erspan_id}",
                        f"origin ip-address {monitoring_ip}",
                        "ip ttl 32",
                        "exit",
                        "no shutdown",
                        "exit",
                    ]
                else:
                    print(
                        f"INFO Site {site} SW{sw_id}: ingen ERSPAN-source opprettet; "
                        "switchen har ingen klient-accessporter å speile."
                    )

        # ERSPAN requires L3 routing on the switch because the MONITORING SVI
        # must be able to reach remote ERSPAN endpoints. Once ip routing is
        # enabled, ip default-gateway is not used for normal routed traffic, so
        # always install a real default route via the MGMT gateway.
        #
        # Do this explicitly from span_mode_site instead of inferring state from
        # the generated command dictionary. This makes the behaviour deterministic
        # and guarantees that SW1/SW2/... get a management return path whenever
        # ERSPAN enables ip routing. More-specific MONITORING /32 routes still win
        # over this default route by longest-prefix match.
        if span_mode_site == "ERSPAN":
            sw_cfg[f"ip route 0.0.0.0 0.0.0.0 {gateway}"] = []
        else:
            sw_cfg[f"ip default-gateway {gateway}"] = []

        sw_cfg[f"ntp server {gateway}"] = []

    return info


def config_vlan(swi_data, site, md, md_top, ip_data, vrf_data, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )

    # IMPORTANT:
    # vlan-antall is a LOCAL access-port requirement, not a statement that the
    # VLAN exists only on that switch.  Build a site-wide union so intermediate
    # switches can actually forward VLANs used farther downstream.
    site_service_vlans = _site_service_vlans(swi_data, md, is_hub)
    dot1x_vlans = _get_dot1x_vlans(md)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)
        local_vlan_info = plan["vlan_info"]
        dot1x_syntax = _switch_compat(
            row, md_top, "dot1x_syntax", ("LEGACY", "AUTH", "ACCESS_SESSION"), "AUTH"
        )

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg["vlan 999"] = ["name NATIVE_UBRUKT", "exit"]

        # Create every site service VLAN on every switch that may need to
        # transport it.  ISP VLAN 200 remains local to HUB SW1 only.
        vlans_to_create = [
            vlan
            for vlan in site_service_vlans
            if vlan != isp_vlan or (is_hub and plan["is_primary_switch"])
        ]
        for vlan in vlans_to_create:
            sw_cfg[f"vlan {vlan}"] = [f"name VLAN_{vlan}", "exit"]

        if span_mode_site == "ERSPAN":
            mon_vlan = monitoring["vlan"]
            service_vlans = {mgmt_vlan, *site_service_vlans}
            if mon_vlan in service_vlans:
                raise ValueError(
                    f"SW{sw_id}: MONITORING VLAN {mon_vlan} kolliderer med et eksisterende tjeneste-VLAN."
                )
            sw_cfg[f"vlan {mon_vlan}"] = ["name MONITORING", "exit"]

        if plan["span_mode"] == "RSPAN":
            rspan_vlan = _get_rspan_vlan(md)
            service_vlans = {mgmt_vlan, *site_service_vlans}
            if rspan_vlan in service_vlans:
                raise ValueError(
                    f"SW{sw_id}: RSPAN-VLAN {rspan_vlan} kolliderer med et tjeneste-VLAN."
                )
            sw_cfg[f"vlan {rspan_vlan}"] = ["name RSPAN_MONITOR", "remote-span", "exit"]

        # LOCAL access-port allocation still comes only from this switch's
        # vlan-antall entry.
        current_port = plan["first_access_port"]
        for vlan, count in local_vlan_info:
            if count <= 0:
                continue

            ports = list(range(current_port, current_port + count))
            current_port += count
            key = _interface_key(intf_prefix, ports)

            port_cfg = [
                f"description access port for VLAN {vlan}",
            ]
            if vlan == isp_vlan:
                # ISP transit port: infrastructure hand-off, not a client access port.
                port_cfg.extend([
                    "switchport mode access",
                    f"switchport access vlan {vlan}",
                    "no shutdown",
                    "exit",
                ])
                sw_cfg[key] = port_cfg
                continue

            port_cfg.extend([
                "switchport mode access",
                f"switchport access vlan {vlan}",
                "switchport port-security",
                "switchport port-security maximum 2",
                "switchport port-security violation restrict",
                "ip verify source",
                "spanning-tree bpduguard enable",
                "spanning-tree portfast",
            ])
            if vlan in dot1x_vlans:
                # Compatibility syntax supported by the older Catalyst 3560
                # and accepted by Catalyst 3850. Do not place 802.1X on
                # infrastructure/server/trunk ports.
                port_cfg.extend([
                    _dot1x_port_control_command(dot1x_syntax),
                    "dot1x pae authenticator",
                ])
            port_cfg.extend(["no shutdown", "exit"])
            sw_cfg[key] = port_cfg

    return info


def _trunk_config(
    vlans, description, channel_group=None, trusted=True, stp_guard=None,
    trunk_encapsulation="AUTO"
):
    """Build trunk config with optional encapsulation command and STP guard."""
    if stp_guard not in {None, "root", "loop"}:
        raise ValueError(f"Ugyldig STP guard: {stp_guard}")

    lines = [description]
    if _include_trunk_encapsulation_command(trunk_encapsulation):
        lines.append("switchport trunk encapsulation dot1q")
    lines.extend([
        "switchport trunk native vlan 999",
        "switchport mode trunk",
        "switchport nonegotiate",
        f"switchport trunk allowed vlan {','.join(map(str, vlans))}",
    ])
    if trusted:
        lines.extend(["ip dhcp snooping trust", "ip arp inspection trust"])
    if channel_group is not None:
        lines.append(f"channel-group {channel_group} mode active")
    if stp_guard is not None:
        lines.append(f"spanning-tree guard {stp_guard}")
    lines.extend(["no shutdown", "exit"])
    return lines


def config_trunk_and_dchp_snooping(swi_data, site, md, md_top, ip_data, vrf_data, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )

    # VLANs required anywhere in the site must traverse the intermediate
    # switch trunks, even when a particular switch has zero local access ports
    # in that VLAN.
    site_service_vlans = _site_service_vlans(swi_data, md, is_hub)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)
        trunk_encapsulation = _switch_compat(
            row, md_top, "trunk_encapsulation",
            ("AUTO", "DOT1Q_CMD", "NO_CMD"), "AUTO"
        )

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        # Do NOT derive trunk VLANs from this switch's vlan-antall.
        # vlan-antall only says how many LOCAL access ports the switch needs.
        all_trunk_vlans = _ordered_unique([mgmt_vlan, *site_service_vlans, 999])
        if span_mode_site == "ERSPAN":
            all_trunk_vlans = _ordered_unique([*all_trunk_vlans[:-1], monitoring["vlan"], 999])

        # VLAN 200 is local ISP transit in this architecture and must not be extended downstream.
        downlink_vlans = [v for v in all_trunk_vlans if v != isp_vlan]

        # RSPAN is a site-local Layer-2 transport VLAN. It is carried only on
        # switch-to-switch trunks, never on the SW1 -> router trunk.
        switch_link_vlans = list(downlink_vlans)
        if plan["span_mode"] == "RSPAN":
            print("ADVARSEL:")
            print("RSPAN samler trafikk fra downstream-switcher.")
            print("Lokal trafikk på RSPAN-destination-switchen SW1 speiles ikke.")
            rspan_vlan = _get_rspan_vlan(md)
            if rspan_vlan in all_trunk_vlans:
                raise ValueError(
                    f"SW{sw_id}: RSPAN-VLAN {rspan_vlan} kolliderer med et eksisterende VLAN."
                )
            switch_link_vlans = _ordered_unique([*switch_link_vlans, rspan_vlan])

        # Root Guard on downlinks assumes SW1 is the intended STP root.
        # Make that role deterministic instead of relying on the lowest MAC.
        if plan["is_primary_switch"] and switch_link_vlans:
            sw_cfg[
                f"spanning-tree vlan {','.join(map(str, switch_link_vlans))} root primary"
            ] = []

        # No snooping/DAI on ISP transit, RSPAN or blackhole/native VLAN 999.
        excluded_inspection_vlans = {isp_vlan, 999}
        if span_mode_site == "ERSPAN":
            excluded_inspection_vlans.add(monitoring["vlan"])
        inspection_vlans = [v for v in all_trunk_vlans if v not in excluded_inspection_vlans]
        if inspection_vlans:
            sw_cfg["ip dhcp snooping"] = []
            sw_cfg[f"ip dhcp snooping vlan {','.join(map(str, inspection_vlans))}"] = []
            sw_cfg["no ip dhcp snooping information option"] = []
            sw_cfg[f"ip arp inspection vlan {','.join(map(str, inspection_vlans))}"] = []

        # UPLINK:
        # - SW1: exactly one physical trunk to the site router.
        # - SW2+: EtherChannel-sized uplink toward the upstream switch when
        #   etherchan_num > 1; otherwise a normal physical trunk.
        if plan["uplink_is_etherchannel"]:
            uplink_po = 1
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                switch_link_vlans,
                f"description UPLINK EtherChannel member(s) - Port-channel{uplink_po}",
                channel_group=uplink_po,
                trusted=True,
                trunk_encapsulation=trunk_encapsulation,
            )
            sw_cfg[f"interface Port-channel{uplink_po}"] = _trunk_config(
                switch_link_vlans,
                f"description UPLINK Port-channel{uplink_po} toward upstream switch "
                f"for VLAN {','.join(map(str, switch_link_vlans))}",
                trusted=True,
                stp_guard="loop",
                trunk_encapsulation=trunk_encapsulation,
            )
            first_downlink_channel = 2
        else:
            uplink_vlans = all_trunk_vlans if plan["is_primary_switch"] else switch_link_vlans
            uplink_desc = (
                "UPLINK trunk to site router"
                if plan["is_primary_switch"]
                else "UPLINK trunk toward upstream switch"
            )
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                uplink_vlans,
                f"description {uplink_desc} for VLAN {','.join(map(str, uplink_vlans))}",
                trusted=True,
                stp_guard=None if plan["is_primary_switch"] else "loop",
                trunk_encapsulation=trunk_encapsulation,
            )
            first_downlink_channel = 1

        # DOWNLINKS:
        # Inter-switch links are trusted for both DHCP snooping and DAI in both
        # directions. Client/access ports remain untrusted by default. This is
        # required because infrastructure devices can use static IP addresses and
        # therefore have no DHCP-snooping binding for DAI validation.
        # Channel-group numbers are local. SW1 starts with Po1. On downstream
        # switches Po1 is reserved for the uplink, so their downlinks start at Po2.
        if plan["downlink_is_etherchannel"]:
            for offset, ports in enumerate(plan["downlink_groups"]):
                channel_id = first_downlink_channel + offset
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK EtherChannel member(s) - Port-channel{channel_id}",
                    channel_group=channel_id,
                    trusted=True,
                    trunk_encapsulation=trunk_encapsulation,
                )
                sw_cfg[f"interface Port-channel{channel_id}"] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK Port-channel{channel_id} for VLAN "
                    f"{','.join(map(str, switch_link_vlans))}",
                    trusted=True,
                    stp_guard="root",
                    trunk_encapsulation=trunk_encapsulation,
                )
        else:
            for idx, ports in enumerate(plan["downlink_groups"], start=1):
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK trunk {idx} for VLAN "
                    f"{','.join(map(str, switch_link_vlans))}",
                    trusted=True,
                    stp_guard="root",
                    trunk_encapsulation=trunk_encapsulation,
                )

        # New num_ports_tot model: skipped ports are real physical interfaces,
        # so explicitly blackhole/shut them instead of leaving them in VLAN 1.
        # if plan["skipped_physical_ports"]:
        #     sw_cfg[_interface_key(intf_prefix, plan["skipped_physical_ports"])] = [
        #         "description UBRUKT - SKIPPED/RESERVED",
        #         "switchport mode access",
        #         "switchport access vlan 999",
        #         "shutdown",
        #         "exit",
        #     ]

        # Remaining unassigned access slots are blackholed and shut down.
        unused_start = plan["first_access_port"] + plan["allocated_access_ports"]
        unused_end = plan["last_access_port"]
        if unused_start <= unused_end:
            unused_ports = list(range(unused_start, unused_end + 1))
            sw_cfg[_interface_key(intf_prefix, unused_ports)] = [
                "description UBRUKT - BLACKHOLE VLAN",
                "switchport mode access",
                "switchport access vlan 999",
                "shutdown",
                "exit",
            ]

    return info


def update_site_config(data, swi_data, sn, conf):
    for raw_sw_id in swi_data["SW"].dropna().unique():
        sw_id = _switch_id(raw_sw_id)
        sw_name = f"SW{sw_id}-SITE-{sn}"
        data[f"site {sn}"]["config"].setdefault(sw_name, {})
        data[f"site {sn}"]["config"][sw_name].update(conf["config"][sw_name])
    return data


def fetch_site_data(config_file):
    try:
        with open(config_file, "r") as f:
            try:
                return json.load(f)
            except json.JSONDecodeError:
                return {}
    except FileNotFoundError:
        return {}


def _refresh_ssh_acls(data, swi_data, sn):
    """
    Record this site's MGMT subnet and re-sync the SSH-MGMT-ONLY ACL on every
    switch across every site processed so far, so newly added sites are
    automatically permitted everywhere (not just locally).
    """
    mgmt_networks = data.setdefault("_mgmt_networks", {})

    row = swi_data.iloc[0]
    try:
        net = ipaddress.ip_network(f"{row.get('MGMT ip')}/{row.get('mask')}", strict=False)
        mgmt_networks[sn] = [str(net.network_address), str(net.hostmask)]
    except (ValueError, TypeError):
        pass

    if not mgmt_networks:
        return

    acl_lines = [f"permit {network} {wildcard}" for network, wildcard in sorted(mgmt_networks.values())]
    acl_lines.append("exit")

    for site_key, site_val in data.items():
        if site_key == "_mgmt_networks":
            continue
        for sw_cfg in site_val.get("config", {}).values():
            if "ip access-list standard SSH-MGMT-ONLY" in sw_cfg:
                sw_cfg["ip access-list standard SSH-MGMT-ONLY"] = acl_lines


def create_site_sw_config(file, sheet, config_file, erspan_context=None):
    sheet_data = read_sheet(file, sheet)
    data = fetch_site_data(config_file)

    md = sheet_data["md"]
    md_top = sheet_data["md_top"]
    swi_data = sheet_data["swi_data"]
    ip_data = sheet_data.get("ip_data")
    vrf_data = sheet_data.get("vrf_data")

    sn = _site_id(md.iloc[0]["site"])
    is_hub = _is_hub_site(md, md_top, sn)
    data[f"site {sn}"] = {"config": {}}

    vlan_conf = config_vlan(swi_data, sn, md, md_top, ip_data, vrf_data, is_hub)
    data = update_site_config(data, swi_data, sn, vlan_conf)

    global_conf = global_config(md, md_top, swi_data, ip_data, vrf_data, is_hub, erspan_context)
    data = update_site_config(data, swi_data, sn, global_conf)

    trunk_conf = config_trunk_and_dchp_snooping(swi_data, sn, md, md_top, ip_data, vrf_data, is_hub)
    data = update_site_config(data, swi_data, sn, trunk_conf)

    _refresh_ssh_acls(data, swi_data, sn)

    with open(config_file, "w") as f:
        json.dump(data, f, indent=4)

    return data


def config_to_text(data, indent=0):
    lines = []
    prefix = "    " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            lines.append("!")
            lines.append(prefix + key)
            lines.extend(config_to_text(value, indent + 1))
    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                lines.append(prefix + value)
            else:
                lines.extend(config_to_text(value, indent))
    elif isinstance(data, str):
        lines.append(prefix + data)

    return lines




def _ansible_normalize_key(key):
    """Normalize generator-only key prefixes such as ``!\ninterface``."""
    parts = [
        line.strip()
        for line in str(key).replace("\r", "").splitlines()
        if line.strip() and line.strip() != "!"
    ]
    return parts[-1] if parts else ""


def _ansible_merge_value(existing, new_value):
    """Merge duplicate normalized Cisco parent blocks while preserving order."""
    if isinstance(existing, dict) and isinstance(new_value, dict):
        merged = dict(existing)
        for key, value in new_value.items():
            if key in merged:
                merged[key] = _ansible_merge_value(merged[key], value)
            else:
                merged[key] = value
        return merged

    existing_list = existing if isinstance(existing, list) else [existing]
    new_list = new_value if isinstance(new_value, list) else [new_value]
    merged = list(existing_list)
    for item in new_list:
        if isinstance(item, str):
            if item not in merged:
                merged.append(item)
        else:
            merged.append(item)
    return merged


def _ansible_normalize_tree(data):
    """Return a hierarchy suitable for ``cisco.ios.ios_config src=...``."""
    if isinstance(data, dict):
        result = {}
        for raw_key, raw_value in data.items():
            key = _ansible_normalize_key(raw_key)
            if not key or key.lower() == "exit":
                continue

            # SSH must already exist before Ansible can push a full config.
            # RSA generation is therefore kept in the manual bootstrap file,
            # not in the idempotency-oriented full Ansible source file.
            if key.lower().startswith("crypto key generate rsa"):
                continue

            value = _ansible_normalize_tree(raw_value)
            if key in result:
                result[key] = _ansible_merge_value(result[key], value)
            else:
                result[key] = value
        return result

    if isinstance(data, list):
        result = []
        for item in data:
            if isinstance(item, str):
                line = item.strip()
                if not line or line in {"!", "exit"}:
                    continue
                # Keep intentional repeated submode terminators such as
                # exit-af-interface. They may occur more than once in one parent.
                result.append(line)
            else:
                result.append(_ansible_normalize_tree(item))
        return result

    if isinstance(data, str):
        line = data.strip()
        return "" if line in {"", "!", "exit"} else line

    return data


def _ansible_config_to_text(data, indent=0):
    """Render normalized Cisco hierarchy without ``!`` or plain ``exit``."""
    lines = []
    prefix = " " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            if not key:
                continue
            lines.append(prefix + str(key))
            lines.extend(_ansible_config_to_text(value, indent + 1))
    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                if value:
                    lines.append(prefix + value)
            else:
                lines.extend(_ansible_config_to_text(value, indent))
    elif isinstance(data, str) and data:
        lines.append(prefix + data)

    return lines


def _switch_bootstrap_config(config_dict):
    """Build a console-paste bootstrap for MGMT + local SSH/SCP.

    It intentionally does not enable TACACS/RADIUS.  The full generated config
    can move VTY login to centralized AAA after Ansible connectivity works.
    """
    config = _ansible_normalize_tree(config_dict)
    selected = {}

    def add_key(key):
        if key in config:
            selected[key] = config[key]

    # Identity first.
    for key in config:
        if key.lower().startswith("hostname "):
            add_key(key)
            break

    # Find the management SVI (VLAN 10 is the project default; fall back to the
    # first SVI that has a static IP address).
    svi_candidates = []
    for key, value in config.items():
        if not key.lower().startswith("interface vlan") or not isinstance(value, list):
            continue
        address = next(
            (
                cmd.split()[2]
                for cmd in value
                if isinstance(cmd, str)
                and cmd.startswith("ip address ")
                and len(cmd.split()) >= 4
                and cmd.split()[2].lower() != "dhcp"
            ),
            None,
        )
        if not address:
            continue
        try:
            vlan_id = int(key.lower().replace("interface vlan", "").strip())
        except ValueError:
            vlan_id = 99999
        svi_candidates.append((0 if vlan_id == 10 else 1, vlan_id, key))

    mgmt_vlan = None
    if svi_candidates:
        _priority, mgmt_vlan, svi_key = sorted(svi_candidates)[0]
        add_key(f"vlan {mgmt_vlan}")
        add_key(svi_key)

        # Include the dedicated MGMT/server access port when the generator has
        # one. This is convenient for direct console/bootstrap cabling.
        for key, value in config.items():
            if not key.lower().startswith("interface ") or key.lower().startswith("interface vlan"):
                continue
            if not isinstance(value, list):
                continue
            has_mgmt_access = any(
                isinstance(cmd, str) and cmd == f"switchport access vlan {mgmt_vlan}"
                for cmd in value
            )
            dedicated = any(
                isinstance(cmd, str)
                and cmd.lower().startswith("description dedicated")
                for cmd in value
            )
            if has_mgmt_access and dedicated:
                selected[key] = [
                    cmd for cmd in value
                    if isinstance(cmd, str)
                    and (
                        cmd == "switchport mode access"
                        or cmd == f"switchport access vlan {mgmt_vlan}"
                        or cmd == "spanning-tree portfast"
                        or cmd == "no shutdown"
                        or cmd.lower().startswith("description dedicated")
                    )
                ]

    # Preserve the generated management default route/gateway.  In ERSPAN mode
    # the switch uses ip routing; otherwise it uses ip default-gateway.
    if "ip routing" in config:
        add_key("ip routing")
    for key in config:
        low = key.lower()
        if low.startswith("ip default-gateway ") or low.startswith("ip route 0.0.0.0 0.0.0.0 "):
            add_key(key)

    # Local SSH/SCP prerequisites.
    for key in config:
        low = key.lower()
        if (
            low.startswith("ip domain name ")
            or low.startswith("username ")
            or low == "ip ssh version 2"
            or low == "ip scp server enable"
            or low == "ip access-list standard ssh-mgmt-only"
        ):
            add_key(key)

    raw_keys = {
        _ansible_normalize_key(k): k for k in config_dict if _ansible_normalize_key(k)
    }
    rsa_key = next(
        (k for k in raw_keys if k.lower().startswith("crypto key generate rsa")),
        None,
    )
    if rsa_key:
        selected[rsa_key] = []

    # Local login makes bootstrap independent of TACACS availability.
    for key, value in config.items():
        if key.lower().startswith("line vty "):
            vty_lines = []
            if isinstance(value, list):
                for cmd in value:
                    if not isinstance(cmd, str):
                        continue
                    if cmd.startswith("login authentication ") or cmd == "login local":
                        continue
                    vty_lines.append(cmd)
            insert_at = 1 if vty_lines and vty_lines[0].startswith("access-class ") else 0
            vty_lines.insert(insert_at, "login local")
            if "transport input ssh" not in vty_lines:
                vty_lines.append("transport input ssh")
            selected[key] = vty_lines
            break

    return selected


def create_ansible_config_files(data):
    """Generate per-site full Ansible configs plus SSH/SCP bootstrap configs."""
    full_root = "ansibleConfigs"
    bootstrap_root = "ansibleBootstrapConfigs"
    os.makedirs(full_root, exist_ok=True)
    os.makedirs(bootstrap_root, exist_ok=True)

    count = 0
    for site, site_data in data.items():
        if str(site).startswith("_") or not isinstance(site_data, dict):
            continue

        site_slug = str(site).strip().lower().replace(" ", "_")
        full_dir = os.path.join(full_root, site_slug)
        boot_dir = os.path.join(bootstrap_root, site_slug)
        os.makedirs(full_dir, exist_ok=True)
        os.makedirs(boot_dir, exist_ok=True)

        for device_name, config_dict in site_data.get("config", {}).items():
            if not isinstance(config_dict, dict):
                continue
            hostname = next(
                (
                    _ansible_normalize_key(k).split(None, 1)[1]
                    for k in config_dict
                    if _ansible_normalize_key(k).lower().startswith("hostname ")
                ),
                str(device_name),
            )

            full_tree = _ansible_normalize_tree(config_dict)
            full_text = _ansible_config_to_text(full_tree)
            with open(os.path.join(full_dir, f"{hostname}.cfg"), "w", encoding="utf-8") as f:
                f.write("\n".join(full_text).rstrip() + "\n")

            bootstrap_tree = _switch_bootstrap_config(config_dict)
            bootstrap_text = _ansible_config_to_text(bootstrap_tree)
            with open(
                os.path.join(boot_dir, f"{hostname}_SSH_SCP.cfg"),
                "w",
                encoding="utf-8",
            ) as f:
                f.write("\n".join(bootstrap_text).rstrip() + "\n")
            count += 1

    print(
        f"Ansible-ready switch-configer: {full_root}/ | "
        f"SSH/SCP-bootstrap: {bootstrap_root}/ ({count} switch(er))."
    )

def create_or_update_config_files(data):
    output_root = "siteSwitchTextConfigs"
    os.makedirs(output_root, exist_ok=True)

    for site, site_data in data.items():
        if site == "_mgmt_networks":
            continue
        site_dir = f"{output_root}/{site}"
        os.makedirs(site_dir, exist_ok=True)

        for sw_name, config in site_data["config"].items():
            text = config_to_text(config)
            with open(f"{site_dir}/{sw_name}.txt", "w", encoding="utf-8") as f:
                f.write("\n".join(text))

    print(f"Text versjon av switch-configene er lagret i {output_root}/")


def create_freeradius_config(file, sheets):
    """Write a FreeRADIUS client snippet using each switch's MGMT SVI address."""
    records = []
    radius_server = None
    radius_key = None

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        md_top = sheet_data["md_top"]
        site = _site_id(md.iloc[0]["site"])
        this_server = _get_radius_server_ip(md_top, required=False)
        this_key = _get_radius_key(md_top, required=False)
        if this_server is None and this_key is None:
            continue
        if this_server is None or this_key is None:
            raise ValueError(f"Site {site}: ufullstendig RADIUS-konfigurasjon i Excel.")
        if radius_server is None:
            radius_server, radius_key = this_server, this_key
        elif (this_server, this_key) != (radius_server, radius_key):
            raise ValueError("RADIUS-server/key må være identisk på alle sites.")

        for _, row in sheet_data["swi_data"].iterrows():
            sw_id = _switch_id(row["SW"])
            name = f"SW{sw_id}-SITE-{site}"
            mgmt_ip = str(ipaddress.ip_address(str(row["MGMT ip"]).strip()))
            records.append((site, int(sw_id), name, mgmt_ip))

    if radius_server is None:
        return

    records.sort(key=lambda x: (int(x[0]), x[1]))
    os.makedirs("freeradius", exist_ok=True)

    lines = [
        "# Generated by NO_MPLS_SAD",
        "# Append/include this snippet from your FreeRADIUS clients.conf.",
        f"# RADIUS server: {radius_server}",
        "",
    ]
    for _site, _sw_id, name, mgmt_ip in records:
        lines.extend([
            f"client {name} {{",
            f"    ipaddr = {mgmt_ip}",
            f"    secret = {radius_key}",
            f"    shortname = {name}",
            "}",
            "",
        ])

    with open("freeradius/clients_network_switches.conf", "w", encoding="utf-8") as f:
        f.write("\n".join(lines).rstrip() + "\n")

    readme = f"""FreeRADIUS / 802.1X generated configuration
===========================================

RADIUS server IP: {radius_server}
Authentication port: {RADIUS_AUTH_PORT}/UDP
Accounting port: {RADIUS_ACCT_PORT}/UDP

1. Install FreeRADIUS (Ubuntu/Debian):
   sudo apt update && sudo apt install freeradius freeradius-utils

2. Append the contents of clients_network_switches.conf to the active
   FreeRADIUS clients.conf (commonly /etc/freeradius/3.0/clients.conf).

3. Configure users/certificates and an EAP method separately. The generator
   intentionally does not create a default username/password.

4. Validate before restart:
   sudo freeradius -XC
   sudo freeradius -X

The generated client IPs are the switches' MGMT SVI addresses. The switches
source RADIUS packets from their MGMT SVI, so these addresses must match.
"""
    with open("freeradius/README_RADIUS.txt", "w", encoding="utf-8") as f:
        f.write(readme)

    print("FreeRADIUS client-konfig er lagret i freeradius/clients_network_switches.conf")


def create_sw_configs_main(file, config_file="site_switch_config.json"):
    sites_sheets = _site_sheet_names(file)

    # Fail fast before any config files are generated if enabled sites mix
    # SPAN/RSPAN/ERSPAN modes. Blank span_mode is allowed and means OFF.
    _validate_span_modes_for_workbook(file, sites_sheets)
    _validate_management_servers_for_workbook(file, sites_sheets)
    erspan_context = _validate_erspan_architecture_for_workbook(file, sites_sheets)

    # Start each generation from a clean JSON model so removed/renamed sites or
    # stale commands from an earlier workbook cannot survive a new run.
    if os.path.exists(config_file):
        os.remove(config_file)

    data = {}
    for sheet in sites_sheets:
        data = create_site_sw_config(file, sheet, config_file, erspan_context)

    create_or_update_config_files(data)
    create_ansible_config_files(data)
    create_freeradius_config(file, sites_sheets)


def main():
    file = sys.argv[1]
    config_file = "site_switch_config.json" if len(sys.argv) < 3 else sys.argv[2]
    create_sw_configs_main(file, config_file)


if __name__ == "__main__":
    main()