import os
import sys
import json
import pandas as pd
from openpyxl import load_workbook
import ipaddress

DEFAULT_IPSEC_PSK = "CHANGE_ME_DMVPN_KEY"
DEFAULT_ISP_VLAN = 200
SSH_DOMAIN = "lab.local"
FLOW_POLICY_SHEET = "FLOW_POLICY"
FLOW_POLICY_COLUMNS = ["source", "destination", "protocol", "destination_port", "action"]


def read_sheet(filename, sheet):
    df = pd.read_excel(
        filename,
        sheet_name=sheet,
        header=None
    )

    md_start = 0
    md_end = df.iloc[md_start:].isna().all(axis=1).idxmax()
    # Metadata width is intentionally dynamic so new product options (for
    # example dns_servers) can be added in Excel without changing this parser.
    md = df.iloc[md_start:md_end].dropna(axis=1, how="all")
    md.columns = md.iloc[0]
    md = md[1:].reset_index(drop=True)

    ip_data_start = md_end + 1
    blank = df.iloc[ip_data_start:].isna().all(axis=1)
    ip_data_end = blank.idxmax()
    ip_data = df.iloc[ip_data_start:ip_data_end, 0:11]
    ip_data.columns = ip_data.iloc[0]
    ip_data = ip_data[1:].reset_index(drop=True)


    vrf_data_start = ip_data_end + 1
    blank = df.iloc[vrf_data_start:].isna().all(axis=1)
    vrf_data_end = blank.idxmax()
    vrf_data = df.iloc[vrf_data_start:vrf_data_end, 0:4]
    vrf_data.columns = vrf_data.iloc[0]
    vrf_data = vrf_data[1:].reset_index(drop=True)
    
    tunnel_data_start = vrf_data_end + 1
    blank = df.iloc[tunnel_data_start:].isna().all(axis=1)
    tunnel_data_end = blank.idxmax()
    tunnel_data = df.iloc[tunnel_data_start:tunnel_data_end, 0:10]
    tunnel_data.columns = tunnel_data.iloc[0]
    tunnel_data = tunnel_data[1:].reset_index(drop=True)

    # print(md)
    # print(ip_data)
    # print(vrf_data)
    # print(tunnel_data)
    # exit(0)
    
    return {
        "md": md,
        "ip_data": ip_data,
        "vrf_data": vrf_data,
        "tunnel_data": tunnel_data,
    }

    # Removed redundant return statement


def is_true(value):
    """
    Godta True/TRUE/1/yes/ja/x fra sheets.
    """
    if pd.isna(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1

    return str(value).strip().lower() in {"true", "1", "yes", "ja", "y", "x"}


def unique_preserve_order(values):
    """Fjern duplikater uten å endre CLI-rekkefølgen."""
    result = []
    for value in values:
        if value not in result:
            result.append(value)
    return result


def ensure_exit_last(commands):
    """Sørg for at Cisco `exit` alltid er siste kommando i en CLI-blokk."""
    commands = [cmd for cmd in commands if cmd != "exit"]
    return unique_preserve_order(commands) + ["exit"]


def getNetId(ip, mask):
    ip = ip.split(".")
    mask = mask.split(".")
    netid = []
    wild_mask = [255, 255, 255, 255]
    wild = []

    for i in range(4):
        ip_b = int(ip[i])
        ip_m = int(mask[i])
        ip_w = wild_mask[i]

        wild.append(str(ip_m ^ ip_w))
        netid.append(str(ip_b & ip_m))

    return ".".join(netid), ".".join(mask), ".".join(wild)



def _normalise_cell_text(value):
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _site_sheet_names(filename):
    """Return only site sheets; FLOW_POLICY is shared workbook metadata."""
    return [
        name
        for name in load_workbook(filename, read_only=True).sheetnames
        if str(name).strip().upper() != FLOW_POLICY_SHEET
    ]


def read_flow_policy(filename):
    """Read the optional shared FLOW_POLICY sheet.

    Removing the sheet (or leaving it empty) keeps backwards compatibility and
    simply disables generated data-plane ACLs.
    """
    try:
        policy = pd.read_excel(filename, sheet_name=FLOW_POLICY_SHEET, dtype=object)
    except ValueError:
        return pd.DataFrame(columns=FLOW_POLICY_COLUMNS)

    if policy.empty:
        return pd.DataFrame(columns=FLOW_POLICY_COLUMNS)

    policy.columns = [str(col).strip().lower() for col in policy.columns]
    missing = [col for col in FLOW_POLICY_COLUMNS if col not in policy.columns]
    if missing:
        raise ValueError(
            f"{FLOW_POLICY_SHEET} mangler kolonne(r): {', '.join(missing)}"
        )

    policy = policy[FLOW_POLICY_COLUMNS].dropna(how="all").reset_index(drop=True)
    for col in FLOW_POLICY_COLUMNS:
        policy[col] = policy[col].map(_normalise_cell_text)
    return policy


def _metadata_value(md, names):
    if md is None or md.empty:
        return None
    row = md.iloc[0]
    for name in names:
        if name in row.index:
            value = row.get(name)
            if not pd.isna(value) and str(value).strip():
                return str(value).strip()
    return None


def get_isp_vlan(md):
    """Return ISP VLAN from router metadata, with VLAN 200 as fallback.

    Supported Excel column names: ISP_vlan, isp_vlan, ISP-VLAN and ISP VLAN.
    An explicit Excel value always overrides DEFAULT_ISP_VLAN.
    """
    value = _metadata_value(
        md,
        ("ISP_vlan", "isp_vlan", "ISP-VLAN", "ISP VLAN"),
    )

    if value is None:
        return DEFAULT_ISP_VLAN

    try:
        vlan = int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Ugyldig ISP VLAN i Excel: {value}") from exc

    if not 1 <= vlan <= 4094 or vlan == 999:
        raise ValueError(
            f"Ugyldig ISP VLAN {vlan}. Velg VLAN 1-4094, men ikke native/blackhole VLAN 999."
        )

    return vlan




def normalise_dmvpn_phase(value):
    """Normaliser DMVPN-fase fra HUB-arket.

    Godtar 2/3, Phase 2/3 og Fase 2/3. Tom verdi beholder bakoverkompatibilitet
    med tidligere generatorversjoner og tolkes som Phase 3.
    """
    if value is None or pd.isna(value) or str(value).strip() == "":
        return 3

    text = str(value).strip().lower().replace("_", " ").replace("-", " ")
    text = " ".join(text.split())
    aliases = {
        "2": 2,
        "phase 2": 2,
        "phase2": 2,
        "fase 2": 2,
        "fase2": 2,
        "3": 3,
        "phase 3": 3,
        "phase3": 3,
        "fase 3": 3,
        "fase3": 3,
    }
    if text in aliases:
        return aliases[text]

    try:
        numeric = int(float(text))
        if numeric in (2, 3):
            return numeric
    except (TypeError, ValueError):
        pass

    raise ValueError(
        f"Ugyldig DMVPN-fase '{value}'. HUB-arket må bruke 2 eller 3 "
        "(eventuelt 'Phase 2'/'Phase 3')."
    )


def get_dmvpn_phase_from_hub(md):
    """Hent nettverksfelles DMVPN-fase fra metadata på HUB-arket."""
    value = _metadata_value(
        md,
        ("dmvpn_phase", "dmvpn phase", "dmvpn_fase", "dmvpn fase"),
    )
    return normalise_dmvpn_phase(value)


def _normalise_compat_choice(value, field, allowed, default):
    """Normalise a small compatibility selector from Excel metadata."""
    if value is None or pd.isna(value) or str(value).strip() == "":
        return default
    text = str(value).strip().upper().replace("-", "_").replace(" ", "_")
    aliases = {
        "OLD": "OLD", "LEGACY": "OLD", "GAMMEL": "OLD",
        "NEW": "NEW", "NY": "NEW",
        "NAMED": "NAMED", "CLASSIC": "CLASSIC",
        "IKEV1": "IKEV1", "IKE1": "IKEV1", "1": "IKEV1",
        "IKEV2": "IKEV2", "IKE2": "IKEV2", "2": "IKEV2",
        "AUTO": "AUTO", "5": "5", "8": "8", "9": "9",
    }
    text = aliases.get(text, text)
    if text not in allowed:
        raise ValueError(
            f"Ugyldig {field}='{value}'. Gyldige verdier: {', '.join(allowed)}"
        )
    return text


def get_tacacs_syntax(md):
    return _normalise_compat_choice(
        _metadata_value(md, ("tacacs_syntax", "tacacs syntax")),
        "tacacs_syntax", ("OLD", "NEW"), "NEW"
    )


def get_eigrp_syntax(md):
    return _normalise_compat_choice(
        _metadata_value(md, ("eigrp_syntax", "eigrp syntax")),
        "eigrp_syntax", ("CLASSIC", "NAMED"), "NAMED"
    )


def get_vrf_syntax(md):
    return _normalise_compat_choice(
        _metadata_value(md, ("vrf_syntax", "vrf syntax")),
        "vrf_syntax", ("OLD", "NEW"), "OLD"
    )


def get_ike_version(md):
    return _normalise_compat_choice(
        _metadata_value(md, ("ike_version", "ike version")),
        "ike_version", ("IKEV1", "IKEV2"), "IKEV2"
    )


def get_secret_type(md):
    return _normalise_compat_choice(
        _metadata_value(md, ("secret_type", "secret type")),
        "secret_type", ("AUTO", "5", "8", "9"), "AUTO"
    )


def _secret_command(prefix, value, secret_type):
    """Build a Cisco secret command without silently relabelling a hash."""
    value = "" if pd.isna(value) else str(value).strip()
    if not value:
        raise ValueError(f"Tom secret/passordverdi for '{prefix}'.")

    detected = None
    for stype in ("5", "8", "9"):
        if value.startswith(f"${stype}$"):
            detected = stype
            break

    if secret_type == "AUTO":
        if detected:
            return f"{prefix} secret {detected} {value}"
        # Plaintext fallback: let IOS hash it according to the platform default.
        return f"{prefix} secret 0 {value}"

    if detected and detected != secret_type:
        raise ValueError(
            f"secret_type={secret_type}, men verdien for '{prefix}' er et type {detected}-hash. "
            "Bytt enten secret_type eller bruk en hash av riktig type."
        )
    if detected is None:
        raise ValueError(
            f"secret_type={secret_type} krever en ferdig type {secret_type}-hash for '{prefix}'. "
            "Bruk AUTO dersom cellen inneholder plaintext."
        )
    return f"{prefix} secret {secret_type} {value}"


def _vrf_forwarding_command(vrf, vrf_syntax):
    return f"vrf forwarding {vrf}" if vrf_syntax == "NEW" else f"ip vrf forwarding {vrf}"

def _append_unique_vrf_endpoint(context, bucket_name, vrf, endpoint):
    """Append a unique endpoint to one VRF-specific context bucket."""
    bucket = context[bucket_name].setdefault(vrf, [])
    if endpoint not in bucket:
        bucket.append(endpoint)


def build_flow_policy_context(filename, site_sheets):
    """Build the workbook-wide address map used by FLOW_POLICY.

    LAN networks and router loopbacks are deliberately kept separate:
      * A VRF name used as a FLOW_POLICY destination means user/LAN networks.
      * Router loopbacks are infrastructure addresses and are NOT implicitly
        included by a VRF-to-VRF policy rule.
      * Loopbacks are still known to the compiler so INTERNET cannot silently
        become a permit to internal router infrastructure.
    """
    context = {
        "vrf_networks": {},
        "vrf_loopbacks": {},
        "dns": [],
        "tacacs": [],
        "syslog": [],
        "radius": [],
    }

    for sheet in site_sheets:
        sheet_data = read_sheet(filename, sheet)
        ip_data = sheet_data["ip_data"]
        vrf_data = sheet_data["vrf_data"]
        md = sheet_data["md"]

        for _, row in ip_data.iterrows():
            vrf = _normalise_cell_text(row.get("vrf")).upper()
            if not vrf:
                continue
            try:
                network = ipaddress.ip_network(
                    f"{row.get('nett id')}/{row.get('mask')}", strict=False
                )
            except ValueError as exc:
                raise ValueError(
                    f"{sheet}: ugyldig nett i VRF {vrf}: "
                    f"{row.get('nett id')} {row.get('mask')}"
                ) from exc
            _append_unique_vrf_endpoint(
                context,
                "vrf_networks",
                vrf,
                ("network", str(network.network_address), str(network.hostmask)),
            )

        for _, row in vrf_data.iterrows():
            vrf = _normalise_cell_text(row.get("vrf")).upper()
            laddr = _normalise_cell_text(row.get("laddr"))
            if vrf and laddr:
                try:
                    laddr = str(ipaddress.ip_address(laddr))
                except ValueError as exc:
                    raise ValueError(f"{sheet}: ugyldig loopback-adresse {laddr}") from exc
                _append_unique_vrf_endpoint(
                    context, "vrf_loopbacks", vrf, ("host", laddr)
                )

        dns_value = _metadata_value(md, ["dns_servers"])
        if dns_value:
            for server in dns_value.replace(",", " ").replace(";", " ").split():
                server = str(ipaddress.ip_address(server))
                if server not in context["dns"]:
                    context["dns"].append(server)

        tacacs = _metadata_value(
            md, ["tacacs_server_ip", "tacacs server ip", "tacacs_server"]
        )
        if tacacs:
            tacacs = str(ipaddress.ip_address(tacacs))
            if tacacs not in context["tacacs"]:
                context["tacacs"].append(tacacs)

        syslog = _metadata_value(
            md,
            ["syslog_server_ip", "rsyslog_server_ip", "syslog server ip", "syslog_server"],
        )
        if syslog:
            syslog = str(ipaddress.ip_address(syslog))
            if syslog not in context["syslog"]:
                context["syslog"].append(syslog)

        radius = _metadata_value(
            md, ["radius_server_ip", "radius server ip", "radius_server"]
        )
        if radius:
            radius = str(ipaddress.ip_address(radius))
            if radius not in context["radius"]:
                context["radius"].append(radius)

    return context


def _get_urpf_command(md):
    """Optional LAN-side uRPF. Default OFF because DHCP/asymmetric designs need care."""
    value = _metadata_value(md, ["urpf_mode", "urpf mode", "uRPF mode"])
    if value is None:
        return None

    mode = value.strip().lower()
    if mode in {"", "off", "none", "disabled", "false", "0"}:
        return None
    if mode in {"strict", "rx"}:
        return "ip verify unicast source reachable-via rx"
    if mode in {"loose", "any"}:
        return "ip verify unicast source reachable-via any"
    raise ValueError("urpf_mode må være OFF, STRICT eller LOOSE")


def _parse_acl_ports(value, protocol):
    value = _normalise_cell_text(value).lower()
    if value in {"", "any", "*"}:
        return [None]

    if protocol not in {"tcp", "udp"}:
        raise ValueError(
            f"destination_port '{value}' kan bare brukes med TCP eller UDP"
        )

    allowed_names = {
        "ssh", "www", "http", "https", "domain", "dns", "tacacs", "syslog",
        "ntp", "bootps", "bootpc"
    }
    result = []
    tokens = value.replace(";", ",").replace(" ", ",").split(",")
    for token in (x.strip() for x in tokens):
        if not token:
            continue
        if "-" in token:
            start, end = token.split("-", 1)
            if not (start.isdigit() and end.isdigit()):
                raise ValueError(f"Ugyldig port-range '{token}'")
            start_i, end_i = int(start), int(end)
            if not (1 <= start_i <= end_i <= 65535):
                raise ValueError(f"Ugyldig port-range '{token}'")
            result.append(("range", str(start_i), str(end_i)))
        elif token.isdigit():
            port = int(token)
            if not 1 <= port <= 65535:
                raise ValueError(f"Ugyldig port '{token}'")
            result.append(("eq", str(port)))
        elif token in allowed_names:
            # IOS accepts common service names in extended ACLs.
            result.append(("eq", "domain" if token == "dns" else token))
        else:
            raise ValueError(f"Ukjent ACL-port/service '{token}'")

    return result or [None]


def _known_flow_vrfs(context):
    return set(context.get("vrf_networks", {})) | set(context.get("vrf_loopbacks", {}))


def _acl_destination_endpoints(
    destination,
    context,
    local_gateway=None,
    source_network=None,
):
    """Resolve one FLOW_POLICY destination.

    Normal service VRFs resolve to LAN/subnet destinations only. Router
    loopbacks are infrastructure addresses and are therefore excluded from
    normal VRF-to-VRF rules.

    MGMT is the exception: MGMT loopbacks are valid management-plane
    destinations (for example SSH to a router loopback), so destination=MGMT
    resolves to both routed MGMT LANs and MGMT loopbacks.

    If source and destination are the same VRF, the local source subnet is
    intentionally kept. Normal host-to-host traffic in the same subnet stays
    at Layer 2 and never reaches this ACL, but traffic addressed to the local
    router gateway does. Keeping the local subnet therefore allows an explicit
    same-VRF FLOW_POLICY rule (for example MGMT -> MGMT permit ip) to cover the
    gateway as expected.
    """
    dest = destination.upper()

    if dest in {"ANY", "INTERNET"}:
        return ["any"]
    if dest == "DNS":
        values = context.get("dns", [])
        if not values:
            raise ValueError("FLOW_POLICY bruker DNS, men dns_servers er ikke satt i Excel")
        return [f"host {ip}" for ip in values]
    if dest == "TACACS":
        values = context.get("tacacs", [])
        if not values:
            raise ValueError("FLOW_POLICY bruker TACACS, men tacacs_server_ip mangler")
        return [f"host {ip}" for ip in values]
    if dest == "SYSLOG":
        values = context.get("syslog", [])
        if not values:
            raise ValueError("FLOW_POLICY bruker SYSLOG, men syslog_server_ip mangler")
        return [f"host {ip}" for ip in values]
    if dest == "RADIUS":
        values = context.get("radius", [])
        if not values:
            raise ValueError("FLOW_POLICY bruker RADIUS, men radius_server_ip mangler")
        return [f"host {ip}" for ip in values]
    if dest == "SELF":
        if not local_gateway:
            raise ValueError("FLOW_POLICY destination=SELF krever lokal gateway")
        return [f"host {local_gateway}"]

    known = _known_flow_vrfs(context)
    if dest not in known:
        raise ValueError(
            f"Ukjent FLOW_POLICY-destination '{destination}'. "
            f"Kjente VRF-er: {', '.join(sorted(known))}; spesialverdier: "
            "INTERNET, ANY, DNS, DHCP, TACACS, SYSLOG, RADIUS, SELF"
        )

    source_network = str(source_network) if source_network is not None else None
    rendered = []
    for endpoint in context.get("vrf_networks", {}).get(dest, []):
        network_address = endpoint[1]
        wildcard = endpoint[2]
        # Do not remove the local source subnet for same-VRF rules.  Although
        # ordinary same-subnet host traffic is switched locally, packets sent to
        # the router's own gateway address do hit this inbound ACL.  A rule such
        # as MGMT -> MGMT permit ip must therefore render the local subnet too.
        rendered.append(f"{network_address} {wildcard}")

    # MGMT loopbacks are intentional management-plane destinations. Keep them
    # addressable by whatever protocol/port the FLOW_POLICY row specifies.
    # Example: MGMT -> MGMT, tcp/22 permits SSH to remote router MGMT loopbacks.
    if dest == "MGMT":
        for endpoint in context.get("vrf_loopbacks", {}).get("MGMT", []):
            value = f"host {endpoint[1]}"
            if value not in rendered:
                rendered.append(value)

    return rendered


def _internet_internal_guards(context, source_vrf, source_network, local_gateway):
    """Return internal destinations that INTERNET must never mean.

    INTERNET is compiled as `any` only after these deny guards.  The guards are
    limited to the source VRF because those are the internal destinations that
    are normally routable from that VRF.  Explicit FLOW_POLICY exceptions are
    emitted before these guards, so an intentional internal permit still wins.
    """
    rendered = []

    if local_gateway:
        rendered.append(f"host {local_gateway}")

    source_network = str(source_network)
    for endpoint in context.get("vrf_networks", {}).get(source_vrf, []):
        network_address = endpoint[1]
        wildcard = endpoint[2]
        mask_int = ((1 << 32) - 1) ^ int(ipaddress.IPv4Address(wildcard))
        candidate = ipaddress.ip_network(
            f"{network_address}/{ipaddress.IPv4Address(mask_int)}", strict=False
        )
        if str(candidate) == source_network:
            continue
        value = f"{network_address} {wildcard}"
        if value not in rendered:
            rendered.append(value)

    for endpoint in context.get("vrf_loopbacks", {}).get(source_vrf, []):
        value = f"host {endpoint[1]}"
        if value not in rendered:
            rendered.append(value)

    return rendered


def _interface_keys_for_vrf(ip_data, intf_prefix, vrf_name):
    intf_nums = list(ip_data["interface"])
    keys = []
    rows = ip_data[ip_data["vrf"].astype(str).str.strip().str.upper() == vrf_name]
    for _, row in rows.iterrows():
        intf = row["interface"]
        vlan = row["vlan"]
        sub = intf_nums.count(intf) > 1
        if sub:
            keys.append(f"interface {intf_prefix}{intf}.{vlan}")
        else:
            keys.append(f"interface {intf_prefix}{intf}")
    return keys


def _insert_before_exit(lines, command):
    if command in lines:
        return
    try:
        idx = lines.index("exit")
    except ValueError:
        lines.append(command)
    else:
        lines.insert(idx, command)


def apply_flow_policy(config, network_info, ip_data, intf_prefix, policy, context):
    """Compile FLOW_POLICY to one ingress extended ACL per local source VRF.

    Specific internal/service rules are emitted before INTERNET rules. INTERNET
    then gets an automatic internal guard so it means "external destination",
    not simply "any routable address".
    """
    if policy is None or policy.empty:
        return

    valid_protocols = {"ip", "tcp", "udp", "icmp", "gre", "esp", "ah"}
    valid_actions = {"permit", "deny"}
    local_vrfs = {
        _normalise_cell_text(v).upper()
        for v in ip_data["vrf"].tolist()
        if _normalise_cell_text(v)
    }

    known_vrfs = _known_flow_vrfs(context)
    for source in policy["source"].tolist():
        source_u = _normalise_cell_text(source).upper()
        if source_u and source_u not in known_vrfs:
            raise ValueError(
                f"Ukjent FLOW_POLICY source '{source}'. Kjente VRF-er: "
                f"{', '.join(sorted(known_vrfs))}"
            )

    # Build ACL definitions separately from the live interface bindings.
    #
    # This is important for Ansible/SSH deployment: if an inbound ACL is bound
    # to the MGMT interface before the ACL has been fully populated, IOS applies
    # the implicit deny immediately. That can terminate the SSH session that is
    # currently pushing the configuration. We therefore render every FLOW ACL
    # before the first interface to which any of them is attached.
    flow_info = {}
    acl_definitions = {}
    acl_bound_interface_keys = []

    for source in policy["source"].tolist():
        source_vrf = _normalise_cell_text(source).upper()
        if not source_vrf or source_vrf not in local_vrfs or source_vrf in flow_info:
            continue

        source_rows = policy[
            policy["source"].astype(str).str.strip().str.upper() == source_vrf
        ]
        ip_rows = ip_data[
            ip_data["vrf"].astype(str).str.strip().str.upper() == source_vrf
        ]
        if ip_rows.empty:
            continue

        source_row = ip_rows.iloc[0]
        source_net = ipaddress.ip_network(
            f"{source_row['nett id']}/{source_row['mask']}", strict=False
        )
        source_match = f"{source_net.network_address} {source_net.hostmask}"
        gateway = _normalise_cell_text(source_row.get("address min"))

        acl_name = f"FLOW-{source_vrf}-IN"
        acl_lines = []
        seen_lines = set()
        seen_rules = {}

        def append_line(line):
            if line not in seen_lines:
                acl_lines.append(line)
                seen_lines.add(line)

        # Validate rules and conflicts before reordering INTERNET behind specific
        # destinations. This keeps the Excel policy deterministic and catches
        # mistakes early.
        row_records = []
        for row_index, row in source_rows.iterrows():
            destination = _normalise_cell_text(row.get("destination")).upper()
            protocol = _normalise_cell_text(row.get("protocol")).lower()
            action = _normalise_cell_text(row.get("action")).lower()
            port_text = _normalise_cell_text(row.get("destination_port")).lower() or "any"

            if not destination:
                raise ValueError(f"FLOW_POLICY rad {row_index + 2}: destination mangler")
            if protocol not in valid_protocols:
                raise ValueError(
                    f"FLOW_POLICY rad {row_index + 2}: ugyldig protocol '{protocol}'"
                )
            if action not in valid_actions:
                raise ValueError(
                    f"FLOW_POLICY rad {row_index + 2}: action må være permit eller deny"
                )

            conflict_key = (source_vrf, destination, protocol, port_text)
            previous = seen_rules.get(conflict_key)
            if previous and previous != action:
                raise ValueError(
                    f"FLOW_POLICY har konflikt for {conflict_key}: både {previous} og {action}"
                )
            if previous == action:
                continue
            seen_rules[conflict_key] = action
            row_records.append((row_index, destination, protocol, action, port_text))

        specific_records = [r for r in row_records if r[1] != "INTERNET"]
        internet_records = [r for r in row_records if r[1] == "INTERNET"]

        def compile_record(record):
            row_index, destination, protocol, action, port_text = record

            if destination == "DHCP":
                if protocol != "udp":
                    raise ValueError("FLOW_POLICY destination=DHCP må bruke protocol=udp")
                if port_text not in {"67", "bootps", "any", "*", ""}:
                    raise ValueError("FLOW_POLICY destination=DHCP forventer destination_port=67")
                append_line(f"{action} udp any eq 68 any eq 67")
                return

            destinations = _acl_destination_endpoints(
                destination,
                context,
                local_gateway=gateway,
                source_network=source_net,
            )
            port_specs = _parse_acl_ports(port_text, protocol)

            # Same-VRF rules include the local LAN as well as remote LANs.
            # Local host-to-host packets stay at Layer 2, while packets to the
            # router gateway are evaluated by this ACL.
            for dest_match in destinations:
                for spec in port_specs:
                    line = f"{action} {protocol} {source_match} {dest_match}"
                    if spec:
                        if spec[0] == "eq":
                            line += f" eq {spec[1]}"
                        else:
                            line += f" range {spec[1]} {spec[2]}"
                    append_line(line)

        # Explicit services and VRF-to-VRF exceptions come first.
        for record in specific_records:
            compile_record(record)

        # INTERNET is not synonymous with ANY. Before external permits, block
        # the router gateway, same-VRF remote LANs and same-VRF loopbacks. Any
        # explicit internal permit above remains effective because ACLs are
        # first-match.
        if any(record[3] == "permit" for record in internet_records):
            for internal_dest in _internet_internal_guards(
                context, source_vrf, source_net, gateway
            ):
                append_line(f"deny ip {source_match} {internal_dest}")

        for record in internet_records:
            compile_record(record)

        if not acl_lines:
            continue

        acl_lines.extend(["deny ip any any log", "exit"])
        # Do not insert the ACL into the main config dict yet. The interface
        # blocks already exist earlier in insertion order, so inserting it here
        # would render `ip access-group ... in` before the ACL is complete.
        acl_definitions[f"ip access-list extended {acl_name}"] = acl_lines

        bound_interfaces = []
        for interface_key in _interface_keys_for_vrf(ip_data, intf_prefix, source_vrf):
            if interface_key not in config:
                raise ValueError(
                    f"Kunne ikke binde {acl_name}: finner ikke {interface_key} i generert config"
                )
            _insert_before_exit(config[interface_key], f"ip access-group {acl_name} in")
            if interface_key not in acl_bound_interface_keys:
                acl_bound_interface_keys.append(interface_key)
            bound_interfaces.append(interface_key.removeprefix("interface "))

        flow_info[source_vrf] = {
            "acl": acl_name,
            "interfaces": bound_interfaces,
            "rules": len(acl_lines) - 2,
            "default": "deny ip any any log",
        }

    # Preserve the rest of the generated configuration order, but place all
    # FLOW ACL definitions immediately before the earliest interface that uses
    # one of them. This guarantees the ACL is complete before IOS receives
    # `ip access-group ... in`, including on the MGMT interface carrying Ansible.
    if acl_definitions:
        current_items = list(config.items())
        bound_key_set = set(acl_bound_interface_keys)
        reordered = {}
        inserted = False

        for key, value in current_items:
            if not inserted and key in bound_key_set:
                reordered.update(acl_definitions)
                inserted = True
            reordered[key] = value

        if not inserted:
            reordered.update(acl_definitions)

        config.clear()
        config.update(reordered)

    if flow_info:
        network_info["flow_policy"] = flow_info

def create_vrf(vrf_data, sn, vrf_syntax="OLD"):
    my_data = {"config": {}, "network_info": {}}

    for _, row in vrf_data.iterrows():
        vrf_name = row["vrf"]
        vrf_loopback = row["loopback"]
        vrf_laddr = row["laddr"]

        my_data["network_info"][vrf_name] = {
            "loopback": vrf_loopback,
            "laddr": vrf_laddr,
        }

        if vrf_syntax == "NEW":
            my_data["config"][f"vrf definition {vrf_name}"] = [
                "address-family ipv4",
                "exit-address-family",
                "exit",
            ]
        else:
            my_data["config"][f"ip vrf {vrf_name}"] = ["exit"]

        my_data["config"][f"interface loopback{vrf_loopback}"] = [
            _vrf_forwarding_command(vrf_name, vrf_syntax),
            f"ip address {vrf_laddr} 255.255.255.255",
            "exit",
        ]

    return my_data


def _allocate_qos_percentages(ip_data, total_percent=75):
    """Allocate integer CBWFQ percentages that sum exactly to total_percent.

    Floors are used first; leftover percentage points go to the largest
    fractional remainders, with higher priority weight winning ties.
    """
    weights = []
    for idx, row in ip_data.iterrows():
        try:
            weight = float(row.get("pri-1-10", 0) or 0)
        except (TypeError, ValueError):
            weight = 0.0
        weights.append((idx, max(weight, 0.0)))

    total_weight = sum(weight for _, weight in weights)
    if total_weight <= 0:
        return {idx: 0 for idx, _ in weights}

    exact = {idx: (weight / total_weight) * total_percent for idx, weight in weights}
    allocated = {idx: int(value) for idx, value in exact.items()}
    remaining = total_percent - sum(allocated.values())

    order = sorted(
        weights,
        key=lambda item: (exact[item[0]] - int(exact[item[0]]), item[1]),
        reverse=True,
    )
    for idx, _weight in order[:remaining]:
        allocated[idx] += 1

    return allocated


def create_interface(ip_data, intf_prefix, md=None):
    my_data = {}
    my_data["config"] = {}
    my_data["network_info"] = {}
    vrf_syntax = get_vrf_syntax(md) if md is not None else "OLD"
    intf_nums = list(ip_data["interface"])
    qos_percentages = _allocate_qos_percentages(ip_data, total_percent=75)
    
    pol_maps= {}

    for index, row in ip_data.iterrows():
        vrf = row["vrf"]
        vlan = row["vlan"]
        pri_afxx = row["pri-afxx"]
        pri_num = row["pri-1-10"]
        top_num = str(pri_afxx)[0]

        pri_prc = qos_percentages.get(index, 0)

        intf = row["interface"]
        sub = True if intf_nums.count(intf) > 1 else False

        ip_address = row["address min"]
        mask = row["mask"]

        if "interfaces" not in my_data["network_info"]:
            my_data["network_info"]["interfaces"] = {}

        my_data["network_info"]["interfaces"][
            f"{intf_prefix}{intf}.{vlan}" if sub else f"{intf_prefix}{intf}"
        ] = {
            "vrf": vrf,
            "vlan": vlan,
            "interface": intf,
            "sub": sub,
            "address": ip_address,
            "mask": mask,
            "pri_afxx": pri_afxx,
            "pri_num": pri_num,
        }

        intf_s = []

        if sub:
            intf_s.append(f"encapsulation dot1Q {vlan}")

            if f"interface {intf_prefix}{intf}.{999}" not in my_data["config"]:
                my_data["config"][
                    f"interface {intf_prefix}{intf}.{999}" if sub
                    else f"interface {intf_prefix}{intf}"
                ] = [
                    "description Sub-interface for ubrukt natiiv VLAN",
                    "encapsulation dot1Q 999 native",
                    "no ip address",
                    "no shutdown",
                    "exit"
                ]

        intf_s.append(_vrf_forwarding_command(vrf, vrf_syntax))
        intf_s.append(f"ip address {ip_address} {mask}")
        intf_s.append("no ip redirects")
        intf_s.append("no ip proxy-arp")
        urpf_command = _get_urpf_command(md)
        if urpf_command:
            intf_s.append(urpf_command)
        intf_s.append("no shutdown")
        intf_s.append("exit")

        my_data["config"][
            f"interface {intf_prefix}{intf}.{vlan}" if sub
            else f"interface {intf_prefix}{intf}"
        ] = intf_s


        if sub:
            my_data["config"][f"class-map match-any QRS-MARK-{vrf}"] = [
                f"match vlan {vlan}",
                "exit"
            ]

            my_data["config"][f"class-map match-any QRS-{vrf}"] = [
                f"match dscp af{pri_afxx}",
                "exit"
            ]

            if f"policy-map QRS-SITE-MARK-POLICY" not in pol_maps:
                pol_maps[f"policy-map QRS-SITE-MARK-POLICY"] = []

            if f"policy-map QRS-SITE-POLICY" not in pol_maps:
                pol_maps[f"policy-map QRS-SITE-POLICY"] = []


            pol_maps[f"policy-map QRS-SITE-MARK-POLICY"].append(
                {
                    f"class QRS-MARK-{vrf}": [
                        f"set dscp af{pri_afxx}", 
                        "exit"
                    ]
                }   
            )

            pol_maps[f"policy-map QRS-SITE-POLICY"].append(
                {
                    f"class QRS-{vrf}": [
                        f"bandwidth percent {pri_prc}", 
                        "exit"
                    ]
                }   
            )
            

    intf_s = []
    intf_s.append("no shutdown")
    intf_s.append("exit")
    my_data["config"][f"interface {intf_prefix}{intf}"] = intf_s
    
    if f"policy-map QRS-SITE-MARK-POLICY" in pol_maps:
        my_data["config"].update(pol_maps)
        
        my_data["config"][f"\ninterface {intf_prefix}{intf}"] = [
            "service-policy input QRS-SITE-MARK-POLICY",
            "exit"
        ]

    if f"policy-map QRS-SITE-POLICY" in pol_maps:
        my_data["config"].update(pol_maps)
        
        my_data["config"][f"\ninterface {intf_prefix}1"] = [
            "service-policy output QRS-SITE-POLICY",
            "exit"
        ]

    return my_data


def create_ipsec_config(
    tunnel, source, sn, sites_data, network_id, vrf, is_hub,
    psk=DEFAULT_IPSEC_PSK, ike_version="IKEV2"
):
    """Generate IPsec for one DMVPN cloud using IKEv1 or IKEv2."""
    suffix = str(network_id).strip()
    transform_set = f"DMVPN-TS-{suffix}"
    ipsec_profile = f"DMVPN-IPSEC-{suffix}"
    config = {}

    other_sites = sites_data.copy()
    other_sites.pop("hub", None)
    other_sites.pop(f"site {sn}", None)

    if ike_version == "IKEV1":
        # IKEv1 is kept deliberately conservative for older IOS platforms.
        # Unique tunnel source-loopbacks let us bind a PSK to each remote peer.
        policy_id = int(float(network_id))
        config[f"crypto isakmp policy {policy_id}"] = [
            "encryption aes 256",
            "hash sha",
            "authentication pre-share",
            "group 14",
            "exit",
        ]

        for site, site_data in other_sites.items():
            tun = site_data.get("network_info", {}).get(tunnel, {})
            peer_source = tun.get("source", "")
            if not peer_source:
                continue
            mask = "255.255.255.255"
            config[f"crypto isakmp key {psk} address {peer_source} {mask}"] = []

            # A newly processed spoke must also be accepted by already-generated
            # peers so Phase 2/3 dynamic spoke-to-spoke IPsec can establish.
            sites_data[site]["config"][
                f"crypto isakmp key {psk} address {source} {mask}"
            ] = []

        config[
            f"crypto ipsec transform-set {transform_set} esp-aes 256 esp-sha-hmac"
        ] = ["mode transport", "exit"]
        config[f"crypto ipsec profile {ipsec_profile}"] = [
            f"set transform-set {transform_set}",
            "exit",
        ]
        return config, ipsec_profile, sites_data

    # IKEv2 / named profile variant.
    proposal = f"DMVPN-IKEV2-PROP-{suffix}"
    policy = f"DMVPN-IKEV2-POL-{suffix}"
    keyring = f"DMVPN-IKEV2-KR-{suffix}"
    ikev2_profile = f"DMVPN-IKEV2-PROFILE-{suffix}"

    config[f"crypto ikev2 proposal {proposal}"] = [
        "encryption aes-cbc-256",
        "integrity sha256",
        "group 14",
        "exit",
    ]
    config[f"crypto ikev2 policy {policy}"] = [f"proposal {proposal}", "exit"]

    remots_s = []
    my_stuff = source
    for site, site_data in other_sites.items():
        tun = site_data.get("network_info", {}).get(tunnel, {})
        ip_address = tun.get("source", "")
        if not ip_address:
            continue
        mask = "255.255.255.255"
        remots_s.append(f"address {ip_address} {mask}")

        keyring_key = f"crypto ikev2 keyring {keyring}"
        profile_key = f"crypto ikev2 profile {ikev2_profile}"
        if keyring_key not in sites_data[site]["config"] or profile_key not in sites_data[site]["config"]:
            raise ValueError(
                f"IPsec-oppsettet for {tunnel} er inkonsistent mellom site {sn} og {site}. "
                "Samme DMVPN-cloud må bruke IKEv2/IPsec på alle deltakende sites."
            )

        site_pers = sites_data[site]["config"][keyring_key][0]["peer ANY"]
        site_pers = [f"address {my_stuff} {mask}"] + site_pers
        sites_data[site]["config"][keyring_key][0]["peer ANY"] = ensure_exit_last(site_pers)

        site_profile_pers = sites_data[site]["config"][profile_key]
        new_match = f"match identity remote address {my_stuff} {mask}"
        site_profile_pers = [new_match] + site_profile_pers
        sites_data[site]["config"][profile_key] = ensure_exit_last(site_profile_pers)

    peer = ensure_exit_last(
        remots_s + [
            f"pre-shared-key local {psk}",
            f"pre-shared-key remote {psk}",
        ]
    )
    config[f"crypto ikev2 keyring {keyring}"] = [
        {"peer ANY": peer},
        "exit",
    ]

    prof_peer = ensure_exit_last(
        [f"match identity remote {remote}" for remote in remots_s]
        + [
            "authentication remote pre-share",
            "authentication local pre-share",
            f"keyring local {keyring}",
        ]
    )
    config[f"crypto ikev2 profile {ikev2_profile}"] = prof_peer

    config[
        f"crypto ipsec transform-set {transform_set} esp-aes 256 esp-sha256-hmac"
    ] = ["mode transport", "exit"]
    config[f"crypto ipsec profile {ipsec_profile}"] = [
        f"set transform-set {transform_set}",
        f"set ikev2-profile {ikev2_profile}",
        "exit",
    ]
    return config, ipsec_profile, sites_data


def create_tunnel_config(tunnel_data, sites_data: dict, is_hub: bool, sn, dmvpn_phase=3, vrf_syntax="OLD", eigrp_syntax="NAMED", ike_version="IKEV2"):
    my_data = {}
    my_data["network_info"] = {}
    my_data["config"] = {}

    if is_hub:
        hub_data = {}
    else:
        hub = sites_data["hub"]
        hub_data = sites_data[hub]["network_info"]

    for idx, row in tunnel_data.iterrows():
        tunnel = row["tunnel id"]
        mode = row["gre mode"]
        ip_address = row["ip address"]
        mask = row["mask"]
        vrf = row["vrf"]
        source = row["source"]
        network_id = row["network-id"]

        ipsec_enabled = is_true(row.get("ipsec", False))

        # Valgfri egen PSK per tunnel.
        # om kollone tom brukes DEFAULT_IPSEC_PSK.
        psk = row.get("ipsec key", DEFAULT_IPSEC_PSK)
        if pd.isna(psk) or str(psk).strip() == "":
            psk = DEFAULT_IPSEC_PSK+f"-{network_id}"
        else:
            psk = str(psk).strip()

        if mode != "multipoint" and "destination" in row.index:
            destination = row["destination"]

        source = str(ipaddress.ip_address(source) + network_id+1)
        my_data["config"][f"interface loopback{network_id+1}"] = [
            f"description Loopback interface for tunnel/ipsec {tunnel}",
            f"ip address {source} 255.255.255.255",
            "ip ospf 1 area 0",
            "no shutdown",
            "exit",
        ]

        tunnel_info = {
            "is hub": is_hub,
            "dmvpn phase": dmvpn_phase,
            "vrf": vrf,
            "ip address": ip_address,
            "mask": mask,
            "source": source,
            "tunnel id": tunnel,
            "mode": mode,
            "ipsec": ipsec_enabled,
        }

        my_data["network_info"][tunnel] = tunnel_info

        tun_s = []
        tun_s.append(_vrf_forwarding_command(vrf, vrf_syntax))
        tun_s.append(f"qos pre-classify")
        tun_s.append(f"ip address {ip_address} {mask}")
        tun_s.append(f"tunnel source loopback{network_id+1}")

        if mode == "multipoint":
            tun_s.append(f"tunnel mode gre {mode}")

            if not is_hub:
                hub_tun_ip = hub_data.get(tunnel, {}).get("ip address", "")
                hub_source = hub_data.get(tunnel, {}).get("source", "")

                tun_s.append(f"ip nhrp map {hub_tun_ip} {hub_source}")
                tun_s.append(f"ip nhrp map multicast {hub_source}")
                tun_s.append(f"ip nhrp nhs {hub_tun_ip}")
            else:
                tun_s.append("ip nhrp map multicast dynamic")

            tun_s.append(f"ip nhrp network-id {network_id}")
            # Phase 3 bruker NHRP Redirect/Shortcut. Phase 2 utelater disse
            # kommandoene og er dermed kompatibel med eldre plattformer som
            # ikke støtter Phase 3.
            if dmvpn_phase == 3:
                if is_hub:
                    tun_s.append("ip nhrp redirect")
                else:
                    tun_s.append("ip nhrp shortcut")

            # In classic EIGRP, interface-specific DMVPN knobs live directly
            # on the tunnel interface instead of under named-mode af-interface.
            if eigrp_syntax == "CLASSIC" and is_hub:
                tun_s.append(f"no ip split-horizon eigrp {int(float(network_id))}")
                if dmvpn_phase == 2:
                    tun_s.append(f"no ip next-hop-self eigrp {int(float(network_id))}")

            tun_s.append("no ip redirects")
            tun_s.append(f"tunnel key {network_id}")

        else:
            print(f"Mode må være multipoint for tunnel {tunnel}")

        # IPsec aktiveres bare når kolonnen 'ipsec' er TRUE/1/yes/ja/x.
        if ipsec_enabled:
            ipsec_config, ipsec_profile, sites_data = create_ipsec_config(
                tunnel, source, sn, sites_data, network_id, vrf, is_hub, psk, ike_version
            )
            my_data["config"].update(ipsec_config)
            tun_s.append(f"tunnel protection ipsec profile {ipsec_profile}")

        tun_s.append("exit")

        my_data["config"][f"interface {tunnel}"] = tun_s
        my_data["network_info"][tunnel] = tunnel_info

    return my_data, sites_data


def create_tunnel_eigrp_config(
    vrf_data, tunnel_data, ip_data, is_hub, dmvpn_phase=3, eigrp_syntax="NAMED"
):
    my_data = {"network_info": {}, "config": {}}
    vrfs = {}

    for _, row in tunnel_data.iterrows():
        network_id = int(float(row["network-id"]))
        tun = row["tunnel id"]
        vrf = row["vrf"]
        ip = row["ip address"]
        mask = row["mask"]
        vrfs[vrf] = {
            "as": network_id,
            "tunnel": tun,
            "networks": [getNetId(ip, mask)],
        }

    for _, row in ip_data.iterrows():
        vrf = row["vrf"]
        if vrf in vrfs:
            vrfs[vrf]["networks"].append(getNetId(row["nett id"], row["mask"]))

    for vrf, data in vrfs.items():
        vrf_laddr = vrf_data[vrf_data["vrf"] == vrf].iloc[0].get("laddr", "")
        if vrf_laddr:
            data["networks"].append((vrf_laddr, "255.255.255.255", "0.0.0.0"))

    if eigrp_syntax == "CLASSIC":
        # Classic mode keeps per-interface split-horizon/next-hop-self commands
        # on the tunnel itself (added by create_tunnel_config).
        router_body = {}
        for vrf, data in vrfs.items():
            af_lines = []
            for network, _mask, wild in data["networks"]:
                af_lines.append(f"network {network} {wild}")
            af_lines.extend([
                "passive-interface default",
                f"no passive-interface {data['tunnel']}",
            ])
            if is_hub and vrf == "INET":
                af_lines.append("redistribute static metric 100000 10 255 1 1500")
            af_lines.append("exit-address-family")
            router_body[
                f"address-family ipv4 vrf {vrf} autonomous-system {data['as']}"
            ] = af_lines
        router_body["exit"] = []
        my_data["config"]["router eigrp 1"] = router_body
        return my_data

    # Named EIGRP mode.
    router_body = {}
    for vrf, data in vrfs.items():
        af_lines = []
        for network, _mask, wild in data["networks"]:
            af_lines.append(f"network {network} {wild}")

        af_lines.extend([
            "\n",
            "af-interface default",
            "passive-interface",
            "exit-af-interface\n",
            f"af-interface {data['tunnel']}",
        ])
        if is_hub:
            af_lines.append("no split-horizon")
            if dmvpn_phase == 2:
                af_lines.append("no next-hop-self")
        af_lines.extend([
            "no passive-interface",
            "exit-af-interface\n",
        ])
        if is_hub and vrf == "INET":
            af_lines.append({
                "topology base": [
                    "redistribute static metric 100000 10 255 1 1500",
                    "exit-af-topology",
                ]
            })
        af_lines.append("exit-address-family")
        router_body[
            f"address-family ipv4 vrf {vrf} autonomous-system {data['as']}"
        ] = af_lines

    router_body["exit"] = []
    my_data["config"]["router eigrp DMVPN-EIGRP"] = router_body
    return my_data


def _get_management_server_ip(md, names, label):
    """Read a management service IP directly from Excel top metadata."""
    if md is None or md.empty:
        raise ValueError(f"{label}-server mangler i Excel-metadata.")

    row = md.iloc[0]
    value = None
    for name in names:
        if name in row.index:
            candidate = row.get(name)
            if not pd.isna(candidate) and str(candidate).strip():
                value = candidate
                break

    if value is None:
        raise ValueError(f"{label}-server mangler i Excel. Forventet felt: {names[0]}.")

    try:
        return str(ipaddress.ip_address(str(value).strip()))
    except ValueError as exc:
        raise ValueError(f"Ugyldig {label}-server-IP i Excel: {value}") from exc


def _get_snmpv3_settings(md):
    """Return validated SNMPv3 polling settings, or None when disabled.

    Excel fields:
      snmpv3_enabled, snmp_server_ip, snmp_user,
      snmp_auth_password, snmp_priv_password

    SHA authentication + AES-128 privacy are intentionally fixed for a
    conservative SNMPv3 authPriv profile that works across the lab platforms.
    """
    if md is None or md.empty:
        return None

    row = md.iloc[0]
    enabled_value = None
    for name in ("snmpv3_enabled", "snmpv3 enabled", "snmp_enabled", "snmp enabled"):
        if name in row.index:
            enabled_value = row.get(name)
            break

    if not is_true(enabled_value):
        return None

    server_ip = _get_management_server_ip(
        md,
        ["snmp_server_ip", "snmp server ip", "nms_server_ip", "nms server ip"],
        "SNMP/NMS",
    )
    user = _metadata_value(md, ("snmp_user", "snmp user"))
    auth_password = _metadata_value(md, ("snmp_auth_password", "snmp auth password"))
    priv_password = _metadata_value(md, ("snmp_priv_password", "snmp priv password"))

    missing = []
    if not user:
        missing.append("snmp_user")
    if not auth_password:
        missing.append("snmp_auth_password")
    if not priv_password:
        missing.append("snmp_priv_password")
    if missing:
        raise ValueError(
            "SNMPv3 er aktivert, men følgende Excel-felt mangler: " + ", ".join(missing)
        )

    if any(ch.isspace() for ch in user):
        raise ValueError("snmp_user kan ikke inneholde mellomrom")
    if len(auth_password) < 8:
        raise ValueError("snmp_auth_password må være minst 8 tegn")
    if len(priv_password) < 8:
        raise ValueError("snmp_priv_password må være minst 8 tegn")

    return {
        "server_ip": server_ip,
        "user": user,
        "auth_password": auth_password,
        "priv_password": priv_password,
    }


def create_snmpv3_config(md):
    """Generate read-only SNMPv3 authPriv polling configuration.

    Access is restricted to the configured NMS IP. Traps are deliberately not
    enabled here; this option is for NMS polling over UDP/161.
    """
    my_data = {"config": {}, "network_info": {}}
    settings = _get_snmpv3_settings(md)
    if settings is None:
        return my_data

    server_ip = settings["server_ip"]
    user = settings["user"]
    auth_password = settings["auth_password"]
    priv_password = settings["priv_password"]

    my_data["config"]["ip access-list standard SNMP-NMS-ONLY"] = [
        f"permit host {server_ip}",
        "deny any",
        "exit",
    ]
    my_data["config"]["snmp-server view NMS-READ iso included"] = []
    my_data["config"]["snmp-server group NMS v3 priv read NMS-READ access SNMP-NMS-ONLY"] = []
    my_data["config"][
        f"snmp-server user {user} NMS v3 auth sha {auth_password} priv aes 128 {priv_password}"
    ] = []
    return my_data


def create_tacacs_config(md, ip_data, sites_data, is_hub):
    my_data = {"config": {}, "network_info": {}}
    if md.empty:
        return my_data

    row = md.iloc[0]
    tacacs_server = _get_management_server_ip(
        md,
        ["tacacs_server_ip", "tacacs server ip", "tacacs_server"],
        "TACACS",
    )
    tacacs_key = row.get("tacacs_key", "")
    if pd.isna(tacacs_key) or not str(tacacs_key).strip():
        raise ValueError("TACACS-key mangler i Excel")
    tacacs_key = str(tacacs_key).strip()
    tacacs_syntax = get_tacacs_syntax(md)

    if is_hub:
        print(f"TACACS server: {tacacs_server} ({tacacs_syntax} syntax)")

    my_data["config"]["aaa new-model"] = []
    if tacacs_syntax == "NEW":
        my_data["config"]["tacacs server TACACS-SERVER"] = [
            f"address ipv4 {tacacs_server}",
            f"key {tacacs_key}",
            "exit",
        ]
        group_server = "server name TACACS-SERVER"
    else:
        my_data["config"][f"tacacs-server host {tacacs_server} key {tacacs_key}"] = []
        group_server = f"server {tacacs_server}"

    my_data["config"]["aaa group server tacacs+ TACACS-GROUP"] = [
        group_server,
        "ip vrf forwarding MGMT",
        "ip tacacs source-interface loopback10",
        "exit",
    ]
    my_data["config"]["aaa authentication login default group TACACS-GROUP local"] = []
    my_data["config"]["aaa authorization exec default group TACACS-GROUP local"] = []
    my_data["config"]["aaa accounting exec default start-stop group TACACS-GROUP"] = []
    my_data["config"]["aaa accounting commands 15 default start-stop group TACACS-GROUP"] = []
    my_data["config"]["line console 0"] = [
        "login authentication default",
        "exec-timeout 10 0",
        "logging synchronous",
        "exit",
    ]
    return my_data


def create_rsyslog_config(md, ip_data, sites_data, is_hub):
    my_data = {"config": {}, "network_info": {}}
    if md.empty:
        return my_data

    rsyslog_server = _get_management_server_ip(
        md,
        ["syslog_server_ip", "rsyslog_server_ip", "syslog server ip", "syslog_server"],
        "Syslog",
    )
    if is_hub:
        print(f"Rsyslog server: {rsyslog_server}")

    my_data["config"]["service timestamps log datetime msec show-timezone"] = []
    my_data["config"][f"logging host {rsyslog_server} vrf MGMT transport udp port 514"] = []
    my_data["config"]["logging trap informational"] = []
    my_data["config"]["logging buffered 16384 informational"] = []
    my_data["config"]["logging source-interface loop10 vrf MGMT"] = []
    return my_data


def enable_ssh(md, vrf_data, ip_data, sites_data, sn, domain=SSH_DOMAIN):
    """
    Genererer lokal SSH-konfig
    brukernavn + passord.
    """
    my_data = {}
    my_data["config"] = {}
    my_data["network_info"] = {}

    if md.empty:
        return my_data

    row = md.iloc[0]

    username = row.get("brukernavn", "")
    password = row.get("passord", "")
    vty_lines = row.get("vty_lines", "0-4")

    if pd.isna(username) or pd.isna(password):
        return my_data

    username = str(username).strip()
    password = str(password).strip()

    if not username or not password:
        return my_data

    my_data["config"][f"ip domain name {domain}"] = []
    secret_type = get_secret_type(md)
    my_data["config"][_secret_command(f"username {username} privilege 15", password, secret_type)] = []
    my_data["config"]["crypto key generate rsa general-keys modulus 2048"] = []
    my_data["config"]["ip ssh version 2"] = []
    my_data["config"]["ip scp server enable"] = []
    
    
    this_site = f"site {sn}"
    
    other_sites = sites_data.copy()

    if "hub" in other_sites:
        del other_sites["hub"]
        
    if this_site in other_sites:
        del other_sites[this_site]
    
    my_mgmt_ip_data = ip_data[ip_data["vrf"] == "MGMT"].iloc[0]
    my_network = my_mgmt_ip_data.get("nett id", "")
    my_mask = my_mgmt_ip_data.get("mask", "")
    my_network, my_mask, my_wild = getNetId(my_network, my_mask)
    
    my_vrf = vrf_data[vrf_data["vrf"] == "MGMT"].iloc[0]
    my_vrf_laddr = my_vrf.get("laddr", "")
    

    
    remots_s = []
    remots_s.append(f"permit {my_network} {my_wild}")
    remots_s.append(f"permit {my_vrf_laddr} 0.0.0.0")
    
    for site, site_data in other_sites.items():
        interfaces = site_data["network_info"].get("interfaces", [])
        for intf, intf_data in interfaces.items():
            vrf = intf_data.get("vrf", "")
            if vrf == "MGMT":
                ip_address = intf_data.get("address", "")
                network = str(ipaddress.ip_address(ip_address) - 1)
                mask = intf_data.get("mask", "")
                network, mask, wild = getNetId(network, mask)
                remots_s.append(f"permit {network} {wild}")

                if f"permit {my_network} {my_wild}" not in sites_data[site]["config"][f"ip access-list standard SSH-MGMT-ONLY"]:
                    sites_data[site]["config"][f"ip access-list standard SSH-MGMT-ONLY"].append(f"permit {my_network} {my_wild}")
                
    my_data["config"][f"ip access-list standard SSH-MGMT-ONLY"] = remots_s

    
    my_data["config"][f"line vty {' '.join(x.strip(' ') for x in vty_lines.split('-'))}"] = [
        "access-class SSH-MGMT-ONLY in vrf-also",
        "login authentication default",
        "exec-timeout 10 0",
        "transport input ssh",
        "exit",
    ]

    return my_data, sites_data


def fetch_site_data(config_file, site_number):
    try:
        data = {}
        hub = ""
        is_hub = False

        with open(config_file, "r") as f:
            try:
                data = json.load(f)
                hub = data.get("hub", "")
            except json.JSONDecodeError:
                data = {}
                data["hub"] = f"site {site_number}"
                hub = f"site {site_number}"

    except FileNotFoundError:
        data = {}
        data["hub"] = f"site {site_number}"
        hub = f"site {site_number}"

    return data, hub == f"site {site_number}"


def create_global_config(md, router_id, intf_prefix, sn, is_hub):
    my_data = {}
    my_data["config"] = {}
    my_data["network_info"] = {}

    secret = md.iloc[0].get("secret", "")
    secret_type = get_secret_type(md)

    my_data["config"][f"hostname RS{sn}"] = []
    my_data["config"][_secret_command("enable", secret, secret_type)] = []
    my_data["config"]["service tcp-keepalives-in"] = []
    my_data["config"]["service tcp-keepalives-out"] = []
    my_data["config"]["no ip source-route"] = []
    my_data["config"]["banner motd # Kun autorisert tilgang er tillatt. Aktivitet kan bli logget. #"] = []

    my_data["config"]["interface loopback0"] = []
    my_data["config"]["interface loopback0"].append(
        f"ip address {router_id} 255.255.255.255"
    )
    my_data["config"]["interface loopback0"].append("ip ospf 1 area 0")
    my_data["config"]["interface loopback0"].append("exit")

    ospf_s = []
    ospf_s.append(f"router-id {router_id}")
    ospf_s.append("exit")
    my_data["config"]["router ospf 1"] = ospf_s


    dhcp = is_true(md.iloc[0].get("DHCP", False))
    dhcp_config = {}
    if is_hub and dhcp:
        dhcp_config = {
            "ip dhcp excluded-address 10.0.0.1 10.0.0.10": [],
            "ip dhcp pool CORE": [
                "network 10.0.0.0 255.255.255.0",
                "default-router 10.0.0.1",
                "exit"
            ]
        }
        
        intf_s = {
            f"interface {intf_prefix}1": [
                "ip address 10.0.0.1 255.255.255.0",
                "ip ospf 1 area 0",
                "no ip redirects",
                "no ip proxy-arp",
                "no shutdown",
                "exit"
            ]
        }
        
        for key, value in dhcp_config.items():
            my_data["config"][key] = value
        for key, value in intf_s.items():
            my_data["config"][key] = value
    else:
        intf_s = []
        intf_s.append("ip address dhcp")
        intf_s.append("ip ospf 1 area 0")
        intf_s.append("no ip redirects")
        intf_s.append("no ip proxy-arp")
        intf_s.append("no shutdown")
        intf_s.append("exit")
        my_data["config"][f"interface {intf_prefix}1"] = intf_s
    


    my_data["network_info"]["loopback0"] = {
        "address": router_id,
        "mask": "255.255.255.255"
    }

    return my_data


def get_dns_servers(md):
    """Return validated DNS server IPs from the optional Excel dns_servers field."""
    if md.empty or "dns_servers" not in md.columns:
        return []

    value = md.iloc[0].get("dns_servers", "")
    if pd.isna(value) or not str(value).strip():
        return []

    # Accept spaces, commas, or semicolons in the Excel cell.
    raw = str(value).replace(",", " ").replace(";", " ")
    servers = [item.strip() for item in raw.split() if item.strip()]
    for server in servers:
        try:
            ipaddress.ip_address(server)
        except ValueError as exc:
            raise ValueError(f"Ugyldig DNS-server i Excel: {server}") from exc
    return servers


def set_up_DHCP_for_vrf_lans(ip_data, md):
    my_config = {}
    my_config["config"] = {}
    my_config["network_info"] = {}

    dns_servers = get_dns_servers(md)

    for idx, row in ip_data.iterrows():
        vrf = row["vrf"]
        ip_gw = row["address min"]
        network = row["nett id"]
        mask = row["mask"]
        num_res = int(float(row["antall-res"]))

        if num_res < 0:
            raise ValueError(f"{vrf}: antall-res kan ikke være negativt ({num_res}).")
        # antall-res is a COUNT including the gateway address.  If 40 addresses
        # are reserved starting at .1, the excluded range must end at .40
        # (not .41).
        ip_res_to = (
            str(ipaddress.ip_address(ip_gw) + num_res - 1)
            if num_res > 0
            else None
        )

        pool = [
            f"vrf {vrf}",
            f"network {network} {mask}",
            f"default-router {ip_gw}",
        ]
        if vrf == "INET" and dns_servers:
            pool.append(f"dns-server {' '.join(dns_servers)}")
        pool.append("exit")

        my_config["config"][f"ip dhcp pool DHCP-{vrf}"] = pool

        if ip_res_to is not None:
            my_config["config"][f"ip dhcp excluded-address vrf {vrf} {ip_gw} {ip_res_to}"] = []

        my_config["network_info"][f"DHCP-{vrf}"] = {
            "network": network,
            "mask": mask,
            "default-router": ip_gw,
            "ip_res_to": ip_res_to
        }
        
    return my_config


def config_nat(md, is_hub):
    my_config = {}
    my_config["config"] = {}
    my_config["network_info"] = {}

    if is_hub:
        intf_prefix = md.iloc[0]["intf_prefix"]
        isp_vlan = get_isp_vlan(md)
        isp_interface = f"{intf_prefix}0.{isp_vlan}"

        my_config["config"]["!\ninterface tunnel20"] = [
            "ip nat inside",
            "exit"
        ]

        my_config["config"][f"interface {isp_interface}"] = [
            f"encapsulation dot1Q {isp_vlan}",
            "ip address dhcp",
            "ip nat outside",
            "no shutdown",
            "exit"
        ]

        my_config["config"]["ip access-list standard NAT-INET"] = [
            "permit any",
            "exit"
        ]

        my_config["config"][
            f"ip nat inside source list NAT-INET interface {isp_interface} vrf INET overload"
        ] = []

        my_config["config"][f"!\ninterface {intf_prefix}0.20"] = [
            "ip nat inside",
            "exit"
        ]

        my_config["config"][
            f"ip route vrf INET 0.0.0.0 0.0.0.0 {isp_interface} dhcp"
        ] = []

        my_config["network_info"]["isp"] = {
            "vlan": isp_vlan,
            "interface": isp_interface,
        }

    return my_config


def config_ntp(sites_data, is_hub):
    my_config = {}
    my_config["config"] = {}
    my_config["network_info"] = {}


    if is_hub:
        my_config["config"]["ntp master 8"] = []
    else:
        hub_site = sites_data["hub"]
        hub_info = sites_data[hub_site]
        server_ip = hub_info["network_info"]["loopback0"]["address"]

        my_config["config"][f"ntp server {server_ip}"] = []

    return my_config
    
    
def _router_site_id(value):
    try:
        f = float(value)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _router_is_hub(md, sheet_name):
    row = md.iloc[0]
    site = _router_site_id(row.get("site"))

    for name in ("is_hub", "hub", "HUB"):
        if name in row.index and not pd.isna(row.get(name)):
            return is_true(row.get(name))
    for name in ("role", "site_role", "site role"):
        if name in row.index and not pd.isna(row.get(name)):
            return str(row.get(name)).strip().upper() == "HUB"
    for name in ("hub_site", "hub site"):
        if name in row.index and not pd.isna(row.get(name)):
            return _router_site_id(row.get(name)) == site

    if "HUB" in str(sheet_name).upper():
        return True
    return site == "1"


def _validate_dmvpn_clouds(filename, site_sheets):
    """Validate that every site uses the same addressing/model per DMVPN cloud."""
    clouds = {}
    seen_source_ips = set()
    for sheet in site_sheets:
        data = read_sheet(filename, sheet)
        md = data["md"]
        site = _router_site_id(md.iloc[0]["site"])
        ike_version = get_ike_version(md)
        for _, row in data["tunnel_data"].iterrows():
            tunnel = _normalise_cell_text(row.get("tunnel id")).lower()
            vrf = _normalise_cell_text(row.get("vrf")).upper()
            network_id = int(float(row.get("network-id")))
            mask = _normalise_cell_text(row.get("mask"))
            ip = _normalise_cell_text(row.get("ip address"))
            ipsec = is_true(row.get("ipsec", False))
            psk = row.get("ipsec key", DEFAULT_IPSEC_PSK)
            if pd.isna(psk) or str(psk).strip() == "":
                psk = DEFAULT_IPSEC_PSK + f"-{network_id}"
            else:
                psk = str(psk).strip()
            try:
                net = ipaddress.ip_network(f"{ip}/{mask}", strict=False)
            except ValueError as exc:
                raise ValueError(f"Site {site} {tunnel}: ugyldig tunnel-IP/mask {ip} {mask}") from exc

            base_source = _normalise_cell_text(row.get("source"))
            try:
                source = str(ipaddress.ip_address(base_source) + network_id + 1)
            except ValueError as exc:
                raise ValueError(f"Site {site} {tunnel}: ugyldig tunnel source-base {base_source}") from exc
            if source in seen_source_ips:
                raise ValueError(f"DMVPN source-loopback {source} brukes mer enn én gang.")
            seen_source_ips.add(source)

            signature = (vrf, network_id, str(net), ipsec, ike_version if ipsec else None, psk if ipsec else None)
            if tunnel in clouds and clouds[tunnel]["signature"] != signature:
                prev = clouds[tunnel]
                raise ValueError(
                    f"DMVPN-cloud {tunnel} er inkonsistent: Site {prev['site']} har "
                    f"{prev['signature']}, Site {site} har {signature}. Samme tunnel må bruke "
                    "samme VRF, network-id, tunnel-subnett, IPsec-status, IKE-versjon og PSK på alle sites."
                )
            clouds.setdefault(tunnel, {"signature": signature, "site": site, "ips": set()})
            if ip in clouds[tunnel]["ips"]:
                raise ValueError(f"DMVPN-cloud {tunnel}: tunnel-IP {ip} brukes av flere sites.")
            clouds[tunnel]["ips"].add(ip)


def configure_site(sheet_file, config_file, sheet, flow_policy=None, flow_context=None, dmvpn_phase=3):
    sheet_data = read_sheet(sheet_file, sheet)

    my_data = {}
    my_data["config"] = {}
    my_data["network_info"] = {}

    #HENT NØDVENDIG DATA 
    ip_data = sheet_data["ip_data"]
    vrf_data = sheet_data["vrf_data"]
    tunnel_data = sheet_data["tunnel_data"]
    md = sheet_data["md"]

    router_id = md.iloc[0]["router-id"]
    sn = md.iloc[0]["site"]
    intf_prefix = md.iloc[0]["intf_prefix"]
    vrf_syntax = get_vrf_syntax(md)
    eigrp_syntax = get_eigrp_syntax(md)
    ike_version = get_ike_version(md)

    data, is_hub = fetch_site_data(config_file, sn)
    
    
    #OPPRETT GLOBAL KONFIGURASJON
    my_data = create_global_config(md,router_id, intf_prefix, sn, is_hub)

    #NTP
    d_ntp = config_ntp(data, is_hub)
    my_data["config"].update(d_ntp["config"])
    
    #VRF
    d_vrf = create_vrf(vrf_data, sn, vrf_syntax)
    my_data["config"].update(d_vrf["config"])
    my_data["network_info"].update(d_vrf["network_info"])

    #TACACS
    d_tacacs = create_tacacs_config(md, ip_data, data, is_hub)
    my_data["config"].update(d_tacacs["config"])

    #RSYSLOG
    d_rsyslog = create_rsyslog_config(md, ip_data, data, is_hub)
    my_data["config"].update(d_rsyslog["config"])

    #SNMPv3 / NMS polling
    d_snmp = create_snmpv3_config(md)
    my_data["config"].update(d_snmp["config"])


    #INTERFACE
    d_ip = create_interface(ip_data, intf_prefix, md)
    my_data["config"].update(d_ip["config"])
    my_data["network_info"].update(d_ip["network_info"])

    # FLOW POLICY / DATA-PLANE ACL
    if flow_policy is not None and not flow_policy.empty:
        apply_flow_policy(
            my_data["config"],
            my_data["network_info"],
            ip_data,
            intf_prefix,
            flow_policy,
            flow_context or {},
        )
    
    #DHCP
    d_dhcp = set_up_DHCP_for_vrf_lans(ip_data, md)
    my_data["config"].update(d_dhcp["config"])
    my_data["network_info"].update(d_dhcp["network_info"])

    #TUNNEL
    d_tunnel, data = create_tunnel_config(
        tunnel_data, data, is_hub, sn, dmvpn_phase, vrf_syntax, eigrp_syntax, ike_version
    )
    my_data["config"].update(d_tunnel["config"])
    my_data["network_info"].update(d_tunnel["network_info"])

    #EIGRP
    d_eigrp = create_tunnel_eigrp_config(
        vrf_data, tunnel_data, ip_data, is_hub, dmvpn_phase, eigrp_syntax
    )
    my_data["config"].update(d_eigrp["config"])
    my_data["network_info"].update(d_eigrp["network_info"])

    #NAT
    d_nat = config_nat(md, is_hub)
    my_data["config"].update(d_nat["config"])

    #SSH
    d_ssh, data = enable_ssh(md, vrf_data, ip_data, data, sn)
    my_data["config"].update(d_ssh["config"])   
     
    #INTerFACE PREFIX
    my_data["intf_prefix"] = intf_prefix

    #LAGRE
    data[f"site {sn}"] = my_data

    with open(config_file, "w") as f:
        json.dump(data, f, indent=4)

    return data


def config_to_text(data, indent=0):
    lines = []
    prefix = "    " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            lines.append("!")
            lines.append(prefix + key)
            lines.extend(config_to_text(value, indent + 1))

    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                lines.append(prefix + value)
            else:
                lines.extend(config_to_text(value, indent))

    elif isinstance(data, str):
        lines.append(prefix + data)
        
    lines.append("!")
  
    return lines




def _ansible_normalize_key(key):
    """Normalize generator-only key prefixes such as ``!\ninterface``."""
    parts = [
        line.strip()
        for line in str(key).replace("\r", "").splitlines()
        if line.strip() and line.strip() != "!"
    ]
    return parts[-1] if parts else ""


def _ansible_merge_value(existing, new_value):
    """Merge duplicate normalized Cisco parent blocks while preserving order."""
    if isinstance(existing, dict) and isinstance(new_value, dict):
        merged = dict(existing)
        for key, value in new_value.items():
            if key in merged:
                merged[key] = _ansible_merge_value(merged[key], value)
            else:
                merged[key] = value
        return merged

    existing_list = existing if isinstance(existing, list) else [existing]
    new_list = new_value if isinstance(new_value, list) else [new_value]
    merged = list(existing_list)
    for item in new_list:
        if isinstance(item, str):
            if item not in merged:
                merged.append(item)
        else:
            merged.append(item)
    return merged


def _ansible_normalize_tree(data):
    """Return a hierarchy suitable for ``cisco.ios.ios_config src=...``.

    The human/copy-paste output intentionally contains ``!`` and explicit
    ``exit`` commands.  The Ansible variant instead uses indentation as the
    hierarchy and merges duplicate parent blocks such as repeated interfaces.
    """
    if isinstance(data, dict):
        result = {}
        for raw_key, raw_value in data.items():
            key = _ansible_normalize_key(raw_key)
            if not key or key.lower() == "exit":
                continue

            # RSA key generation is an interactive/bootstrap action.  A device
            # must already have working SSH before Ansible can push this file.
            if key.lower().startswith("crypto key generate rsa"):
                continue

            value = _ansible_normalize_tree(raw_value)
            if key in result:
                result[key] = _ansible_merge_value(result[key], value)
            else:
                result[key] = value
        return result

    if isinstance(data, list):
        result = []
        for item in data:
            if isinstance(item, str):
                line = item.strip()
                if not line or line in {"!", "exit"}:
                    continue
                # Keep intentional repeated submode terminators such as
                # exit-af-interface. They may occur more than once in one parent.
                result.append(line)
            else:
                result.append(_ansible_normalize_tree(item))
        return result

    if isinstance(data, str):
        line = data.strip()
        return "" if line in {"", "!", "exit"} else line

    return data


def _ansible_config_to_text(data, indent=0):
    """Render normalized Cisco hierarchy without ``!`` or plain ``exit``."""
    lines = []
    prefix = " " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            if not key:
                continue
            lines.append(prefix + str(key))
            lines.extend(_ansible_config_to_text(value, indent + 1))
    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                if value:
                    lines.append(prefix + value)
            else:
                lines.extend(_ansible_config_to_text(value, indent))
    elif isinstance(data, str) and data:
        lines.append(prefix + data)

    return lines


def _router_bootstrap_config(site_data):
    """Build the minimum practical console bootstrap for Ansible SSH/SCP.

    This includes the generated MGMT VRF/interface and local SSH credentials,
    but deliberately avoids TACACS/RADIUS so the first Ansible login does not
    depend on external AAA.  WAN/DMVPN reachability is not recreated here.
    """
    config = _ansible_normalize_tree(site_data.get("config", {}))
    mgmt_ip = None
    interfaces = site_data.get("network_info", {}).get("interfaces", {})
    for _name, info in interfaces.items():
        if str(info.get("vrf", "")).upper() == "MGMT" and info.get("address"):
            mgmt_ip = str(info["address"])
            break

    selected = {}

    def add_key(key):
        if key in config:
            selected[key] = config[key]

    # Identity and management VRF.
    for key in config:
        low = key.lower()
        if low.startswith("hostname "):
            add_key(key)
        elif low == "ip vrf mgmt" or low == "vrf definition mgmt":
            add_key(key)

    # Management interface used by the generated Ansible inventory.
    mgmt_parent = None
    if mgmt_ip:
        for key, value in config.items():
            if not key.lower().startswith("interface ") or not isinstance(value, list):
                continue
            if any(
                isinstance(cmd, str)
                and cmd.lower().startswith(f"ip address {mgmt_ip} ")
                for cmd in value
            ):
                mgmt_parent = key
                # Keep only commands needed to establish management reachability.
                # Data-plane ACL/uRPF/QoS are intentionally left for the full
                # Ansible deployment so bootstrap cannot lock out Ansible.
                selected[key] = [
                    cmd for cmd in value
                    if isinstance(cmd, str)
                    and (
                        cmd.startswith("encapsulation dot1Q ")
                        or cmd.lower() in {"ip vrf forwarding mgmt", "vrf forwarding mgmt"}
                        or cmd.lower().startswith(f"ip address {mgmt_ip} ")
                        or cmd == "no shutdown"
                    )
                ]
                break

    # If MGMT lives on a subinterface, make sure the physical parent is up.
    if mgmt_parent and "." in mgmt_parent.split(None, 1)[1]:
        physical = mgmt_parent.split(None, 1)[1].split(".", 1)[0]
        parent_key = f"interface {physical}"
        if parent_key in config:
            parent_lines = [
                cmd for cmd in config[parent_key]
                if isinstance(cmd, str) and cmd in {"no shutdown"}
            ]
            if parent_lines:
                selected[parent_key] = parent_lines

    # Local SSH/SCP bootstrap.  RSA generation stays here because this file is
    # intended to be pasted from console before Ansible is usable.
    raw_config = site_data.get("config", {})
    normalized_raw_keys = {
        _ansible_normalize_key(k): k for k in raw_config if _ansible_normalize_key(k)
    }
    for key in config:
        low = key.lower()
        if (
            low.startswith("ip domain name ")
            or low.startswith("username ")
            or low == "ip ssh version 2"
            or low == "ip scp server enable"
            or low == "ip access-list standard ssh-mgmt-only"
        ):
            add_key(key)

    rsa_key = next(
        (k for k in normalized_raw_keys if k.lower().startswith("crypto key generate rsa")),
        None,
    )
    if rsa_key:
        # Empty list = one top-level command when rendered.
        selected[rsa_key] = []

    # Use local authentication for bootstrap even though the final generated
    # configuration later moves VTY login to the normal AAA policy.
    for key, value in config.items():
        if key.lower().startswith("line vty "):
            vty_lines = []
            if isinstance(value, list):
                for cmd in value:
                    if not isinstance(cmd, str):
                        continue
                    if cmd.startswith("login authentication ") or cmd == "login local":
                        continue
                    vty_lines.append(cmd)
            insert_at = 1 if vty_lines and vty_lines[0].startswith("access-class ") else 0
            vty_lines.insert(insert_at, "login local")
            if "transport input ssh" not in vty_lines:
                vty_lines.append("transport input ssh")
            selected[key] = vty_lines
            break

    return selected


def create_ansible_config_files(data):
    """Generate idempotency-friendly full configs and console bootstrap files."""
    full_root = "ansibleConfigs"
    bootstrap_root = "ansibleBootstrapConfigs"
    os.makedirs(full_root, exist_ok=True)
    os.makedirs(bootstrap_root, exist_ok=True)

    count = 0
    for site, site_data in data.items():
        if site == "hub" or not isinstance(site_data, dict):
            continue

        site_slug = str(site).strip().lower().replace(" ", "_")
        hostname = next(
            (
                _ansible_normalize_key(k).split(None, 1)[1]
                for k in site_data.get("config", {})
                if _ansible_normalize_key(k).lower().startswith("hostname ")
            ),
            site_slug.upper(),
        )

        full_dir = os.path.join(full_root, site_slug)
        boot_dir = os.path.join(bootstrap_root, site_slug)
        os.makedirs(full_dir, exist_ok=True)
        os.makedirs(boot_dir, exist_ok=True)

        full_tree = _ansible_normalize_tree(site_data.get("config", {}))
        full_text = _ansible_config_to_text(full_tree)
        with open(os.path.join(full_dir, f"{hostname}.cfg"), "w", encoding="utf-8") as f:
            f.write("\n".join(full_text).rstrip() + "\n")

        bootstrap_tree = _router_bootstrap_config(site_data)
        bootstrap_text = _ansible_config_to_text(bootstrap_tree)
        with open(
            os.path.join(boot_dir, f"{hostname}_SSH_SCP.cfg"),
            "w",
            encoding="utf-8",
        ) as f:
            f.write("\n".join(bootstrap_text).rstrip() + "\n")
        count += 1

    print(
        f"Ansible-ready router-configer: {full_root}/ | "
        f"SSH/SCP-bootstrap: {bootstrap_root}/ ({count} router(e))."
    )

def create_or_update_config_files(data):
    if not os.path.exists("siteEdgeRouterTextConfigs"):
        os.makedirs("siteEdgeRouterTextConfigs")

    for site, site_data in data.items():
        if site == "hub" or not isinstance(site_data, dict):
            continue
        config = site_data["config"]
        text = config_to_text(config)

        with open(
            f"siteEdgeRouterTextConfigs/EDGE_ROUTER_{site.replace(' ', '_').upper()}.txt",
            "w",
            encoding="utf-8"
        ) as f:
            f.write("\n".join(text))
    
    print()
    print(f"Text-versjoner av edge-router-configene er lagret i siteEdgeRouterTextConfigs/.")
    print()
    

def create_edge_router_configs_main(file, config_file="EDGE_ROUTER_configs.json"):   

    sites_sheets = _site_sheet_names(file)
    if not sites_sheets:
        raise ValueError("Arbeidsboken inneholder ingen site-ark.")

    _validate_dmvpn_clouds(file, sites_sheets)

    # Resolve the HUB explicitly instead of assuming that the first sheet is HUB.
    hub_sheets = []
    for sheet in sites_sheets:
        md = read_sheet(file, sheet)["md"]
        if _router_is_hub(md, sheet):
            hub_sheets.append(sheet)
    if len(hub_sheets) != 1:
        raise ValueError(f"Forventet nøyaktig ett HUB-site, fant {len(hub_sheets)}.")
    hub_sheet = hub_sheets[0]
    hub_md = read_sheet(file, hub_sheet)["md"]
    hub_site = _router_site_id(hub_md.iloc[0]["site"])
    dmvpn_phase = get_dmvpn_phase_from_hub(hub_md)
    print(f"DMVPN Phase {dmvpn_phase} valgt fra HUB-arket ({hub_sheet}).")
    ordered_sheets = [hub_sheet] + [s for s in sites_sheets if s != hub_sheet]

    flow_policy = read_flow_policy(file)
    flow_context = (
        build_flow_policy_context(file, sites_sheets)
        if flow_policy is not None and not flow_policy.empty
        else {}
    )

    # Clean model prevents stale sites/commands from previous runs.
    with open(config_file, "w", encoding="utf-8") as f:
        json.dump({"hub": f"site {hub_site}"}, f, indent=4)

    for sheet in ordered_sheets:
        data = configure_site(
            file,
            config_file,
            sheet,
            flow_policy=flow_policy,
            flow_context=flow_context,
            dmvpn_phase=dmvpn_phase,
        )
    
    create_or_update_config_files(data)
    create_ansible_config_files(data)


def main():
    if not os.path.exists("siteEdgeRouterTextConfigs"):
        os.makedirs("siteEdgeRouterTextConfigs")

    sheet_file = sys.argv[1]
    config_file = (
        sys.argv[2]
        if len(sys.argv) > 2
        else "EDGE_ROUTER_configs.json"
    )

    create_edge_router_configs_main(sheet_file, config_file)


if __name__ == "__main__":
    main()